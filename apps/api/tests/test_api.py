import os
from pathlib import Path

from fastapi.testclient import TestClient
from pydantic import ValidationError

os.environ["DATA_DIR"] = "/tmp/ragbench-test-data"
os.environ["RAGBENCH_PROVIDER_MODE"] = "deterministic"
os.environ["RAGBENCH_QUEUE_MODE"] = "inline"
os.environ["RAGBENCH_VECTOR_STORE"] = "sqlite"
os.environ["RAGBENCH_MEDIA_RUNTIME"] = "container"

from app.database import connect, data_dir, init_db  # noqa: E402
from app import main as main_module  # noqa: E402
from app import agentic as agentic_module  # noqa: E402
from app import media_processing  # noqa: E402
from app.agentic import AGENT_LIMITS, TOOL_NAMES, EvidencePackage, ToolRegistry, AgentState, structured_critic_report, structured_orchestrator_plan  # noqa: E402
from app.contracts import (  # noqa: E402
    GroundingReport,
    PROMPT_REGISTRY,
    RunConfig,
    VisualObservation,
    validate_answer_contract,
)
from app.diagnostics import local_log_path, scrub  # noqa: E402
from app.guardrails import MVP_EXCLUSIONS, REQUIRED_DEFERRED_SCOPE, assert_no_forbidden_defaults  # noqa: E402
from app.hybrid import combined_vector_graph_rerank, extract_query_entities, graph_expand_entities  # noqa: E402
from app.job_execution import (  # noqa: E402
    BENCHMARK_STEPS,
    CeleryJobRunner,
    INGESTION_STEPS,
    JOB_STATES,
    LocalJobRunner,
    RedisRQJobRunner,
    fetch_next_queued_job,
)
from app.migrations import LATEST_SCHEMA_VERSION, migration_status  # noqa: E402
from app.main import LIMITS, app  # noqa: E402
from app.providers import (  # noqa: E402
    DEFAULT_OLLAMA_LLM_MODEL,
    DEFAULT_OLLAMA_VLM_MODEL,
    DeterministicEmbeddingAdapter,
    OllamaAdapter,
    ProviderTimeout,
    ProviderUnavailable,
    ProviderResult,
    TesseractOCRAdapter,
    VisualObservation as ProviderVisualObservation,
    choose_document_parser,
    ollama_generation_payload,
)
from app.provider_extensions import PaidCloudProviderStub, assert_local_first_provider_policy, extension_settings  # noqa: E402
from app.provider_health import provider_health  # noqa: E402
from app.rag import RetrievedChunk, complete_required_fact_mentions, compute_comparison, keyword_retrieve, recommend, rerank_evidence, retrieve, run_pipeline, vector_retrieve  # noqa: E402
from app.reliability import DEFAULT_RERANKED_CHUNKS, DEFAULT_TOP_K_CHUNKS  # noqa: E402
from app.resource_profile import GIB, preflight_report  # noqa: E402
from app.runtime_adapters import (  # noqa: E402
    ChromaVectorStoreAdapter,
    NetworkXGraphStoreAdapter,
    QdrantVectorStoreAdapter,
    SQLiteGraphStoreAdapter,
    SQLiteMetadataRepository,
    equivalent_graph_results,
    equivalent_vector_results,
    normalize_ocr_output,
    normalize_transcript_output,
    vector_metadata,
)
from app.storage import ARTIFACT_LAYOUT, parse_trace_jsonl  # noqa: E402
from app.workflow import GRAPH_NODES, RAGState  # noqa: E402
from app.worker import process_job  # noqa: E402


client = TestClient(app)


def setup_function():
    import shutil

    shutil.rmtree(os.environ["DATA_DIR"], ignore_errors=True)
    main_module.MODEL_SETTINGS_OVERRIDES.clear()
    init_db()


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_sqlite_connections_wait_for_transient_locks():
    with connect() as conn:
        busy_timeout = conn.execute("PRAGMA busy_timeout").fetchone()[0]
        journal_mode = conn.execute("PRAGMA journal_mode").fetchone()[0]

    assert busy_timeout >= 30000
    assert journal_mode.lower() == "wal"


def test_project_upload_and_run_text_vertical_slice():
    project = client.post("/api/projects", json={"name": "Demo"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("requirements.txt", b"Authentication uses API Gateway and Token Store. Deployment has latency risk.", "text/plain")},
    )
    assert upload.status_code == 200
    assert upload.json()["project"]["chunk_count"] == 1

    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What authentication components are mentioned?"},
    )
    assert run.status_code == 200
    body = run.json()
    assert len(body["results"]) == 3
    assert body["recommendation"]["recommended_flow"] in {"traditional", "agentic", "hybrid_graph"}
    assert all(result["citations"] for result in body["results"])
    assert body["knowledge_base_version_id"] == upload.json()["project"]["active_version_id"]
    assert body["traditional_result_id"]
    assert body["agentic_result_id"]
    assert body["hybrid_result_id"]


def test_run_without_chunks_returns_controlled_error():
    project = client.post("/api/projects", json={"name": "Empty"}).json()
    response = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "Anything?"},
    )
    assert response.status_code == 400
    assert "Upload" in response.json()["detail"]


def test_no_matching_evidence_returns_warning():
    project = client.post("/api/projects", json={"name": "Mismatch"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("notes.txt", b"Bananas and oranges are listed in the kitchen inventory.", "text/plain")},
    )
    response = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "How does authentication work?"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["recommendation"]["recommended_flow"] == "insufficient_evidence"
    assert all("insufficient_evidence" in result["warnings"] for result in body["results"])


def test_project_update_and_delete():
    project = client.post("/api/projects", json={"name": "Before", "description": "old"}).json()
    updated = client.patch(
        f"/api/projects/{project['id']}",
        json={"name": "After", "description": "new"},
    )
    assert updated.status_code == 200
    assert updated.json()["name"] == "After"
    assert updated.json()["description"] == "new"

    deleted = client.delete(f"/api/projects/{project['id']}")
    assert deleted.status_code == 200
    assert client.get(f"/api/projects/{project['id']}").status_code == 404


def test_upload_versions_append_clear_replace_and_create_new():
    project = client.post("/api/projects", json={"name": "Versioned"}).json()
    first = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("one.txt", b"Authentication uses tokens.", "text/plain")},
    ).json()
    assert first["project"]["active_version"]["version_number"] == 1
    first_version_id = first["project"]["active_version_id"]

    second = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"upload_action": "append"},
        files={"files": ("two.txt", b"Deployment has latency risk.", "text/plain")},
    ).json()
    assert second["project"]["active_version"]["version_number"] == 2
    assert second["project"]["active_version"]["change_type"] == "append_sources"

    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What uses tokens?"},
    ).json()
    run_version = run["knowledge_base_version_id"]

    rejected = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"upload_action": "clear_replace"},
        files={"files": ("three.txt", b"New clean data.", "text/plain")},
    )
    assert rejected.status_code == 400
    assert "confirm_clear" in rejected.json()["detail"]

    cleared = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"upload_action": "clear_replace", "confirm_clear": "true"},
        files={"files": ("three.txt", b"New clean data.", "text/plain")},
    ).json()
    assert cleared["project"]["source_count"] == 1
    assert cleared["project"]["active_version"]["change_type"] == "clear_and_replace"
    assert client.get(f"/api/runs/{run['id']}").json()["knowledge_base_version_id"] == run_version
    assert run_version != first_version_id

    created = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"upload_action": "create_new", "new_project_name": "Separate"},
        files={"files": ("four.txt", b"Separate knowledge base data.", "text/plain")},
    ).json()
    assert created["project"]["id"] != project["id"]
    assert created["project"]["name"] == "Separate"
    assert created["project"]["active_version"]["change_type"] == "initial_upload"


def test_source_delete_creates_remove_version():
    project = client.post("/api/projects", json={"name": "Delete Source"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("delete-me.txt", b"Temporary source content.", "text/plain")},
    ).json()
    source_id = upload["sources"][0]["id"]

    deleted = client.delete(f"/api/sources/{source_id}")
    assert deleted.status_code == 200
    project_after = deleted.json()["project"]
    assert project_after["source_count"] == 0
    assert project_after["active_version"]["change_type"] == "remove_sources"
    assert client.get(f"/api/sources/{source_id}").status_code == 404


def test_follow_up_run_metadata_keeps_source_evidence_authority():
    project = client.post("/api/projects", json={"name": "Follow Up"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    first = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?", "user_feedback": "good"},
    ).json()
    second = client.post(
        f"/api/projects/{project['id']}/runs",
        json={
            "question": "Can you expand on that?",
            "parent_run_id": first["id"],
            "run_mode": "all_pipelines",
            "selected_sources": [],
        },
    )
    assert second.status_code == 200
    body = second.json()
    assert body["parent_run_id"] == first["id"]
    assert body["knowledge_base_version_id"] == first["knowledge_base_version_id"]

    other_project = client.post("/api/projects", json={"name": "Other"}).json()
    invalid = client.post(
        f"/api/projects/{other_project['id']}/runs",
        json={"question": "Invalid", "parent_run_id": first["id"]},
    )
    assert invalid.status_code in {400, 422}


def test_text_ingestion_creates_content_blocks_chunks_and_vectors():
    project = client.post("/api/projects", json={"name": "Blocks"}).json()
    response = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("notes.md", b"# Heading\nAuthentication uses API Gateway.", "text/markdown")},
    )
    assert response.status_code == 200
    source_id = response.json()["sources"][0]["id"]

    with connect() as conn:
        blocks = conn.execute("SELECT * FROM content_blocks WHERE source_id = ?", (source_id,)).fetchall()
        chunks = conn.execute("SELECT * FROM chunks WHERE source_id = ?", (source_id,)).fetchall()
        vectors = conn.execute("SELECT * FROM vector_records WHERE source_id = ?", (source_id,)).fetchall()

    assert len(blocks) == 1
    assert blocks[0]["block_type"] == "text"
    assert len(chunks) == 1
    assert chunks[0]["content_block_id"] == blocks[0]["id"]
    assert chunks[0]["embedding_id"] == vectors[0]["id"]
    assert chunks[0]["source_reference_label"] == chunks[0]["citation"]
    assert len(vectors) == 1


def test_media_ingestion_creates_derived_blocks_and_artifacts():
    project = client.post("/api/projects", json={"name": "Media"}).json()
    response = client.post(
        f"/api/projects/{project['id']}/sources",
        files=[
            ("files", ("diagram.png", b"not-a-real-image", "image/png")),
            ("files", ("meeting.wav", b"not-a-real-audio", "audio/wav")),
            ("files", ("demo.mp4", b"not-a-real-video", "video/mp4")),
        ],
    )
    assert response.status_code == 200
    assert response.json()["project"]["source_count"] == 3

    with connect() as conn:
        blocks = conn.execute(
            "SELECT block_type, frame_path, metadata_json FROM content_blocks WHERE project_id = ?",
            (project["id"],),
        ).fetchall()
    block_types = [block["block_type"] for block in blocks]
    assert "ocr" in block_types
    assert "caption" in block_types
    assert "transcript" in block_types
    assert any(block["frame_path"] for block in blocks)


def test_partial_success_when_one_file_fails():
    project = client.post("/api/projects", json={"name": "Partial"}).json()
    response = client.post(
        f"/api/projects/{project['id']}/sources",
        files=[
            ("files", ("good.txt", b"Useful source text.", "text/plain")),
            ("files", ("bad.xyz", b"Unsupported source text.", "application/octet-stream")),
        ],
    )
    assert response.status_code == 200
    body = response.json()
    assert body["project"]["processing_status"] == "partial_success"
    statuses = {source["filename"]: source["status"] for source in body["sources"]}
    assert statuses["good.txt"] == "processed"
    assert statuses["bad.xyz"] == "failed"


def test_file_size_limit_rejects_oversized_upload():
    original = LIMITS["max_file_size_mb"]
    LIMITS["max_file_size_mb"] = 0
    try:
        project = client.post("/api/projects", json={"name": "Limit"}).json()
        response = client.post(
            f"/api/projects/{project['id']}/sources",
            files={"files": ("too-big.txt", b"x", "text/plain")},
        )
        assert response.status_code == 200
        body = response.json()
        assert body["sources"][0]["status"] == "failed"
        assert "max_file_size_mb" in body["sources"][0]["error"]
    finally:
        LIMITS["max_file_size_mb"] = original


def test_retrieval_ranks_matching_evidence_above_related_noise():
    project = client.post("/api/projects", json={"name": "Ranking"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files=[
            ("files", ("auth.txt", b"Authentication uses API Gateway and Token Store for login requests.", "text/plain")),
            ("files", ("deploy.txt", b"Deployment notes discuss latency budgets and cache warming.", "text/plain")),
        ],
    )

    hits = retrieve(project["id"], "What authentication components handle login?", top_k=3)

    assert hits
    assert "Authentication uses API Gateway" in hits[0].text
    assert hits[0].rerank_score >= hits[-1].rerank_score
    assert hits[0].keyword_score > 0


def test_vector_keyword_merge_deduplicates_and_preserves_scores():
    project = client.post("/api/projects", json={"name": "Merge"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication tokens are verified by the Token Store.", "text/plain")},
    )

    vector_hits = vector_retrieve(project["id"], "How are authentication tokens verified?", limit=5)
    keyword_hits = keyword_retrieve(project["id"], "How are authentication tokens verified?", limit=5)
    merged_hits = retrieve(project["id"], "How are authentication tokens verified?", top_k=5)

    assert vector_hits
    assert keyword_hits
    assert len({hit.id for hit in merged_hits}) == len(merged_hits)
    assert merged_hits[0].vector_score > 0
    assert merged_hits[0].keyword_score > 0


def test_metadata_filtering_limits_retrieval_to_selected_source_type():
    project = client.post("/api/projects", json={"name": "Filters"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files=[
            ("files", ("diagram.png", b"fake-image", "image/png")),
            ("files", ("notes.txt", b"Diagram review notes mention architecture arrows.", "text/plain")),
        ],
    )

    hits = retrieve(
        project["id"],
        "What does the diagram mention?",
        top_k=5,
        metadata_filters={"source_type": "image"},
    )

    assert hits
    assert {hit.metadata["source_type"] for hit in hits} == {"image"}


def test_vector_metadata_removes_none_and_non_scalar_values_for_external_stores():
    cleaned = vector_metadata(
        {
            "source_id": "source-1",
            "page_number": None,
            "timestamp_start": 0.0,
            "tags": ["image"],
            "nested": {"a": 1},
            "visual_caption": True,
        }
    )

    assert cleaned == {"source_id": "source-1", "timestamp_start": 0.0, "visual_caption": True}


def test_reranking_prefers_dense_grounded_evidence():
    project = client.post("/api/projects", json={"name": "Rerank"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files=[
            ("files", ("dense.txt", b"Authentication token rotation policy requires weekly key review.", "text/plain")),
            ("files", ("sparse.txt", b"Authentication is mentioned briefly. Other content covers release notes.", "text/plain")),
        ],
    )
    candidates = vector_retrieve(project["id"], "authentication token rotation policy", limit=10)

    reranked = rerank_evidence("authentication token rotation policy", candidates, top_k=2)

    assert reranked
    assert "token rotation policy" in reranked[0].text
    assert reranked[0].rerank_score >= reranked[-1].rerank_score


def test_answer_completion_keeps_required_model_names_from_evidence():
    answer = complete_required_fact_mentions(
        "What are the default LLM, VLM, and embedding models?",
        "The product uses host Ollama for real LLM, VLM, and embedding benchmarks.",
        "The default models are qwen3:8b for LLM, qwen3-vl:4b for VLM, and bge-m3 for embeddings.",
    )

    assert "qwen3:8b" in answer
    assert "qwen3-vl:4b" in answer
    assert "bge-m3" in answer


def test_traditional_pipeline_trace_citations_and_grounding_metrics():
    project = client.post("/api/projects", json={"name": "Traditional"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store for user sessions.", "text/plain")},
    )

    result = run_pipeline(project["id"], "What handles user sessions?", "traditional")

    steps = [entry["step"] for entry in result["trace"]]
    assert "retrieve_vector_chunks" in steps
    assert "retrieve_keyword_chunks" in steps
    assert "merge_and_rank_chunks" in steps
    assert "format_citations" in steps
    assert result["citations"]
    assert result["metrics"]["grounding_score"] > 0
    assert result["metrics"]["insufficient_evidence_flag"] is False


def test_selected_source_filter_and_no_evidence_guardrail():
    project = client.post("/api/projects", json={"name": "Selected Sources"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files=[
            ("files", ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")),
            ("files", ("fruit.txt", b"Bananas and oranges are listed in the kitchen inventory.", "text/plain")),
        ],
    ).json()
    source_by_name = {source["filename"]: source["id"] for source in upload["sources"]}

    selected = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?", "selected_sources": [source_by_name["fruit.txt"]]},
    ).json()

    assert selected["recommendation"]["recommended_flow"] == "insufficient_evidence"
    assert all(result["metrics"]["insufficient_evidence_flag"] for result in selected["results"])


def test_agentic_tool_registry_is_limited_to_mvp_tools():
    state = AgentState(project_id="project", question="question")
    registry = ToolRegistry(state)

    assert registry.names == [
        "search_vector_store",
        "search_keyword_index",
        "rerank_evidence",
        "inspect_chunk",
        "inspect_visual_source",
        "refine_visual_question",
        "inspect_transcript",
        "lookup_entities",
        "evaluate_grounding",
        "final_answer",
    ]
    try:
        registry.call("unknown_tool")
    except ValueError as exc:
        assert "not registered" in str(exc)
    assert "malformed_tool_call" in state.warnings


def test_agentic_evidence_package_contract_validates_required_fields():
    package = EvidencePackage(
        agent_name="Retrieval Agent",
        task="Find evidence.",
        query_used="What handles authentication?",
        evidence_items=[],
        citations=[],
        confidence=0.0,
        gaps=["No evidence."],
        conflicts=[],
        recommended_next_action="insufficient_evidence",
    )

    assert package.agent_name == "Retrieval Agent"
    assert package.model_dump()["recommended_next_action"] == "insufficient_evidence"


def test_agentic_pipeline_respects_limits_and_persists_independently():
    project = client.post("/api/projects", json={"name": "Agentic"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store for user sessions.", "text/plain")},
    )

    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What handles user sessions and what is related?"},
    ).json()
    agentic = next(result for result in run["results"] if result["pipeline_type"] == "agentic")
    traditional = next(result for result in run["results"] if result["pipeline_type"] == "traditional")

    assert run["agentic_result_id"] != run["traditional_result_id"]
    assert agentic["citations"]
    assert agentic["metrics"]["agent_steps"] <= AGENT_LIMITS["max_agent_steps"]
    assert agentic["metrics"]["tool_calls"] <= AGENT_LIMITS["max_total_tool_calls"]
    assert agentic["metrics"]["retrieval_calls"] <= AGENT_LIMITS["max_retrieval_calls"]
    assert agentic["metrics"]["graph_agent_calls"] <= AGENT_LIMITS["max_graph_agent_calls"]
    assert agentic["metrics"]["specialist_agents_used"] <= AGENT_LIMITS["max_specialist_agents"]
    assert traditional["pipeline_type"] == "traditional"


def test_agentic_trace_contains_orchestrator_packages_critic_and_final_answer():
    project = client.post("/api/projects", json={"name": "Agentic Trace"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )

    result = run_pipeline(project["id"], "What does authentication use?", "agentic")

    steps = [entry["step"] for entry in result["trace"]]
    assert "orchestrator_plan" in steps
    assert "dispatch_retrieval_agent" in steps
    assert "critic_grounding_check" in steps
    assert "finalize_answer" in steps
    packages = [entry["package"] for entry in result["trace"] if "package" in entry]
    assert any(package["agent_name"] == "Retrieval Agent" for package in packages)
    assert any(package["agent_name"] == "Critic / Grounding Agent" for package in packages)
    assert all(package["agent_name"] != "Main Orchestrator" for package in packages)


def test_agentic_visual_path_returns_controlled_vlm_unavailable_warning():
    project = client.post("/api/projects", json={"name": "Visual Agent"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("architecture.png", b"fake-image", "image/png")},
    )

    result = run_pipeline(project["id"], "What does the architecture diagram show?", "agentic")

    steps = [entry["step"] for entry in result["trace"]]
    assert "dispatch_visual_agent_if_needed" in steps
    assert "vlm_unavailable_used_derived_visual_evidence" in result["warnings"]
    assert result["metrics"]["vlm_calls"] <= AGENT_LIMITS["max_vlm_calls"]


def test_agentic_no_evidence_and_selected_source_guardrail():
    project = client.post("/api/projects", json={"name": "Agentic Guardrail"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("fruit.txt", b"Bananas and oranges are listed in the kitchen inventory.", "text/plain")},
    )

    result = run_pipeline(project["id"], "How does authentication work?", "agentic")

    assert result["answer"].startswith("I could not find enough evidence")
    assert result["citations"] == []
    assert "insufficient_evidence" in result["warnings"]
    assert result["metrics"]["critic_passes"] <= AGENT_LIMITS["max_critic_passes"]


def test_hybrid_extracts_entities_and_expands_graph_evidence():
    project = client.post("/api/projects", json={"name": "Hybrid Graph"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("architecture.txt", b"API Gateway connects to Token Store. Token Store supports Session Service.", "text/plain")},
    )

    entities = extract_query_entities(project["id"], "How is API Gateway related to Token Store?")
    graph_chunks = graph_expand_entities(project["id"], entities)

    assert entities
    assert graph_chunks
    assert graph_chunks[0].metadata["graph_relationship"]["type"] == "co_occurs_with"


def test_hybrid_combined_reranking_applies_graph_boost():
    project = client.post("/api/projects", json={"name": "Hybrid Rerank"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("architecture.txt", b"API Gateway connects to Token Store for authentication.", "text/plain")},
    )
    retrieved = retrieve(project["id"], "How does API Gateway connect to Token Store?", top_k=8)
    graph_chunks = graph_expand_entities(project["id"], extract_query_entities(project["id"], "API Gateway Token Store"))

    combined = combined_vector_graph_rerank("How does API Gateway connect to Token Store?", retrieved, graph_chunks, top_k=8)

    assert combined
    assert any(chunk.metadata.get("graph_boost_applied") for chunk in combined)


def test_hybrid_pipeline_answer_citations_and_trace():
    project = client.post("/api/projects", json={"name": "Hybrid Pipeline"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("architecture.txt", b"API Gateway connects to Token Store for authentication.", "text/plain")},
    )

    result = run_pipeline(project["id"], "How does API Gateway connect to Token Store?", "hybrid_graph")

    steps = [entry["step"] for entry in result["trace"]]
    assert "extract_query_entities" in steps
    assert "graph_expand_entities" in steps
    assert "rerank_combined_evidence" in steps
    assert "verify_grounding" in steps
    assert result["citations"]
    assert result["metrics"]["query_entities_count"] > 0
    assert result["metrics"]["graph_chunks_count"] > 0
    assert result["metrics"]["grounding_score"] > 0


def test_hybrid_visual_check_and_grounding_no_evidence_paths():
    visual_project = client.post("/api/projects", json={"name": "Hybrid Visual"}).json()
    client.post(
        f"/api/projects/{visual_project['id']}/sources",
        files={"files": ("architecture.png", b"fake-image", "image/png")},
    )
    visual_result = run_pipeline(visual_project["id"], "What does the architecture diagram show?", "hybrid_graph")

    assert "targeted_visual_check_if_needed" in [entry["step"] for entry in visual_result["trace"]]
    assert "vlm_unavailable_used_derived_visual_evidence" in visual_result["warnings"]
    assert visual_result["metrics"]["vlm_calls"] <= AGENT_LIMITS["max_vlm_calls"]

    empty_project = client.post("/api/projects", json={"name": "Hybrid No Evidence"}).json()
    client.post(
        f"/api/projects/{empty_project['id']}/sources",
        files={"files": ("fruit.txt", b"Bananas and oranges are listed in the kitchen inventory.", "text/plain")},
    )
    no_evidence = run_pipeline(empty_project["id"], "How does authentication work?", "hybrid_graph")

    assert no_evidence["answer"].startswith("I could not find enough graph-grounded evidence")
    assert "insufficient_evidence" in no_evidence["warnings"]


def test_comparison_normalizes_metrics_and_quality_labels():
    results = [
        {
            "pipeline_type": "traditional",
            "metrics": {"citations_count": 2, "grounding_score": 0.8, "unsupported_claims_count": 0, "latency_seconds": 0.1},
            "warnings": [],
        },
        {
            "pipeline_type": "agentic",
            "metrics": {"citations_count": 0, "grounding_score": 0.0, "unsupported_claims_count": 1, "insufficient_evidence_flag": True},
            "warnings": ["insufficient_evidence"],
        },
    ]

    comparison = compute_comparison(results)

    assert comparison["pipelines"]["traditional"]["quality_label"] == "strong"
    assert comparison["pipelines"]["agentic"]["quality_label"] == "insufficient"
    assert comparison["best_pipeline"] == "traditional"
    assert comparison["pipelines"]["traditional"]["evidence_score"] <= 1


def test_recommendation_prefers_hybrid_when_graph_evidence_is_used():
    results = [
        {
            "pipeline_type": "traditional",
            "metrics": {"citations_count": 1, "grounding_score": 0.5, "unsupported_claims_count": 0},
            "warnings": [],
        },
        {
            "pipeline_type": "agentic",
            "metrics": {"citations_count": 1, "grounding_score": 0.5, "unsupported_claims_count": 0},
            "warnings": [],
        },
        {
            "pipeline_type": "hybrid_graph",
            "metrics": {"citations_count": 1, "grounding_score": 0.5, "unsupported_claims_count": 0, "graph_chunks_count": 1},
            "warnings": [],
        },
    ]

    recommendation = recommend(results)

    assert recommendation["recommended_flow"] == "hybrid_graph"
    assert recommendation["quality_label"] in {"usable", "strong"}
    assert recommendation["tradeoff_notes"]


def test_run_response_and_exports_include_comparison_json_markdown():
    project = client.post("/api/projects", json={"name": "Exports"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()

    assert run["comparison"]["best_pipeline"] in {"traditional", "agentic", "hybrid_graph"}
    assert "quality_label" in run["comparison"]["pipelines"]["traditional"]

    exported_json = client.get(f"/api/runs/{run['id']}/export.json")
    assert exported_json.status_code == 200
    assert exported_json.json()["schema"] == "ragbench.run_export.v1"
    assert exported_json.json()["run"]["comparison"]["tradeoff_notes"]

    exported_markdown = client.get(f"/api/runs/{run['id']}/export.md")
    assert exported_markdown.status_code == 200
    text = exported_markdown.text
    assert "## Comparison" in text
    assert "Tradeoffs:" in text
    assert "Quality labels:" in text


def test_weak_evidence_recommendation_returns_insufficient_not_best_effort():
    project = client.post("/api/projects", json={"name": "Weak Evidence"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("fruit.txt", b"Bananas and oranges are listed in the kitchen inventory.", "text/plain")},
    )

    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "How does authentication work?"},
    ).json()

    assert run["recommendation"]["recommended_flow"] == "insufficient_evidence"
    assert run["comparison"]["best_pipeline"] == "insufficient_evidence"
    assert all(data["quality_label"] == "insufficient" for data in run["comparison"]["pipelines"].values())


def test_dependency_endpoint_reports_missing_tools_without_blocking_core_models():
    response = client.get("/api/settings/dependencies")

    assert response.status_code == 200
    body = response.json()
    assert body["python"]["available"] is True
    assert body["embeddings"]["available"] is True
    assert "ffmpeg" in body
    assert "tesseract" in body
    assert "available" in body["ollama"]
    assert "warning" in body["ollama"]
    if body["ollama"]["available"]:
        assert "models" in body["ollama"]


def test_run_events_and_trace_jsonl_are_persisted_and_parseable():
    project = client.post("/api/projects", json={"name": "Diagnostics"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()

    events = client.get(f"/api/runs/{run['id']}/events")
    assert events.status_code == 200
    event_types = [event["event_type"] for event in events.json()]
    assert "create_run" in event_types
    assert "persist_outputs" in event_types
    assert events.json()[-1]["status"] == "succeeded"

    trace = client.get(f"/api/runs/{run['id']}/trace.jsonl")
    assert trace.status_code == 200
    lines = [line for line in trace.text.splitlines() if line.strip()]
    assert lines
    import json

    parsed = [json.loads(line) for line in lines]
    assert parsed[0]["event"] == "recommendation"
    assert any(item["event"] == "trace_step" for item in parsed)


def test_diagnostic_scrubber_removes_secrets_and_absolute_data_paths():
    secret_payload = {
        "api_key": "abc",
        "nested": {"token": "secret-token"},
        "path": str(local_log_path()),
    }

    cleaned = scrub(secret_payload)

    assert cleaned["api_key"] == "[redacted]"
    assert cleaned["nested"]["token"] == "[redacted]"
    assert "$DATA_DIR" in cleaned["path"]


def test_local_logs_are_written_without_secrets_for_run_events():
    project = client.post("/api/projects", json={"name": "Logs"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway.", "text/plain")},
    )
    client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What uses API Gateway?"},
    )

    log_text = local_log_path().read_text(encoding="utf-8")
    assert "persist_outputs" in log_text
    assert "/tmp/ragbench-test-data" not in log_text
    assert "api_key" not in log_text


def test_diagnostics_log_endpoint_and_frontend_events_are_structured():
    import json

    request_id = "test-request-123"
    health = client.get("/health", headers={"X-Request-ID": request_id})
    frontend = client.post(
        "/api/diagnostics/frontend-event",
        json={"level": "info", "event_type": "playwright_probe", "payload": {"step": "click_run", "api_key": "hidden"}},
    )
    logs = client.get("/api/diagnostics/logs?limit=50")

    assert health.status_code == 200
    assert health.headers["x-request-id"] == request_id
    assert frontend.status_code == 200
    assert logs.status_code == 200
    entries = logs.json()["entries"]
    assert any(entry["event_type"] == "api_request_completed" and entry["payload"].get("request_id") == request_id for entry in entries)
    assert any(entry["event_type"] == "frontend_event" and entry["payload"].get("event_type") == "playwright_probe" for entry in entries)
    assert "hidden" not in json.dumps(entries)


def test_provider_status_endpoint_exposes_defaults_and_adapter_choices():
    response = client.get("/api/settings/providers")

    assert response.status_code == 200
    body = response.json()
    assert body["defaults"]["ollama"]["llm_model"] == DEFAULT_OLLAMA_LLM_MODEL
    assert body["defaults"]["ollama"]["vlm_model"] == DEFAULT_OLLAMA_VLM_MODEL
    assert "ollama pull" in body["defaults"]["ollama"]["setup"][1]
    assert body["embedding_adapters"]["deterministic"]["available"] is True
    assert "tesseract" in body["ocr_adapters"]
    assert "whisper_cpp" in body["transcription_adapters"]


def test_ollama_adapter_reports_missing_model_with_setup_guidance():
    def fake_transport(url, payload, timeout):
        assert url.endswith("/api/tags")
        return {"models": [{"name": "other-model"}]}

    adapter = OllamaAdapter(model="missing-model", transport=fake_transport)
    validation = adapter.validate_model()

    assert validation["available"] is False
    assert validation["setup"] == "ollama pull missing-model"
    try:
        adapter.generate("hello")
    except ProviderUnavailable as exc:
        assert "missing-model" in str(exc)
    else:
        raise AssertionError("Expected ProviderUnavailable")


def test_ollama_adapter_accepts_default_latest_tag():
    def fake_transport(url, payload, timeout):
        assert url.endswith("/api/tags")
        return {"models": [{"name": "bge-m3:latest"}]}

    adapter = OllamaAdapter(model="bge-m3", transport=fake_transport)

    assert adapter.validate_model()["available"] is True


def test_ollama_generation_disables_thinking_by_default():
    payload = ollama_generation_payload("qwen3:8b", "answer directly")

    assert payload["think"] is False
    assert payload["stream"] is False


def test_ollama_adapter_sends_final_answer_payload_and_warns_on_empty_thinking_response():
    calls = []

    def fake_transport(url, payload, timeout):
        calls.append((url, payload))
        if url.endswith("/api/tags"):
            return {"models": [{"name": "qwen3:8b"}]}
        return {"response": "", "thinking": "internal reasoning", "done_reason": "length"}

    adapter = OllamaAdapter(model="qwen3:8b", transport=fake_transport)
    result = adapter.generate("hello")

    assert calls[-1][1]["think"] is False
    assert result.text == ""
    assert result.warnings
    assert result.metadata["done_reason"] == "length"


def test_ollama_adapter_timeout_path_is_controlled():
    def timeout_transport(url, payload, timeout):
        raise ProviderTimeout("timed out")

    adapter = OllamaAdapter(model="model", transport=timeout_transport)
    try:
        adapter.validate_model()
    except ProviderTimeout as exc:
        assert "timed out" in str(exc)
    else:
        raise AssertionError("Expected ProviderTimeout")


def test_embedding_and_parser_contracts_are_deterministic():
    embeddings = DeterministicEmbeddingAdapter()
    one = embeddings.embed_text("Authentication uses Token Store")
    batch = embeddings.embed_batch(["Authentication uses Token Store", "Other"])

    assert len(one) == 16
    assert batch[0] == one
    assert choose_document_parser("file.pdf")["parser"] == "pypdf"
    assert choose_document_parser("file.docx")["parser"] == "python-docx"
    assert choose_document_parser("file.pdf", prefer_docling=True)["parser"] in {"docling", "fallback"}


def test_ocr_adapter_missing_dependency_returns_clear_error():
    adapter = TesseractOCRAdapter()
    status = adapter.availability()

    assert "available" in status
    if not status["available"]:
        try:
            adapter.extract_text(local_log_path())
        except ProviderUnavailable as exc:
            assert "Tesseract" in str(exc)
        else:
            raise AssertionError("Expected ProviderUnavailable when Tesseract is absent")


def test_epic11_graph_nodes_and_state_limits_are_inspectable():
    response = client.get("/api/settings/graph-nodes")

    assert response.status_code == 200
    assert response.json()["nodes"]["traditional"] == GRAPH_NODES["traditional"]
    state = RAGState(pipeline_type="traditional", question="What is cited?")
    assert state.can_retry_revision() is True
    state.revision_attempts = 1
    assert state.can_retry_revision() is False
    state.transition("validate", "validated")
    state.transition("unknown", "bad")
    assert state.completed_nodes == ["validate"]
    assert state.errors[0]["error"] == "unknown_node"


def test_epic11_completed_run_persists_reconstructable_graph_states():
    project = client.post("/api/projects", json={"name": "Graph State"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()

    states = client.get(f"/api/runs/{run['id']}/graph-states")

    assert states.status_code == 200
    body = states.json()
    assert {item["pipeline_type"] for item in body} == {"traditional", "agentic", "hybrid_graph"}
    traditional = [item for item in body if item["pipeline_type"] == "traditional"][-1]
    assert traditional["state"]["completed_nodes"] == GRAPH_NODES["traditional"]
    assert traditional["state"]["metrics"]["graph_runtime"] == "langgraph"
    assert traditional["state"]["limits"]["max_answer_revision_attempts"] == 1
    assert traditional["state"]["final_answer"]


def test_epic12_prompt_registry_and_contract_schemas_validate():
    prompts = client.get("/api/settings/prompts")
    assert prompts.status_code == 200
    assert set(PROMPT_REGISTRY).issubset(prompts.json()["prompts"])

    observation = VisualObservation(
        source_id="source",
        question="What is visible?",
        answer="A login diagram is visible.",
        visible_text=["Login"],
        described_objects=["diagram"],
        confidence=0.8,
    )
    assert observation.needs_follow_up is False

    report = GroundingReport(
        supported_claims=["Authentication uses API Gateway."],
        unsupported_claims=[],
        missing_evidence=[],
        conflicting_evidence=[],
        citation_quality="cited",
        answer_sufficient=True,
        recommended_action="finalize",
    )
    assert report.answer_sufficient is True

    try:
        GroundingReport(citation_quality="bad", answer_sufficient=False, recommended_action="invent")
    except ValidationError:
        pass
    else:
        raise AssertionError("Expected invalid recommended action to fail validation")


def test_epic12_run_config_and_prompt_versions_are_persisted_on_reopen():
    project = client.post("/api/projects", json={"name": "Run Config"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    ).json()
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()

    reopened = client.get(f"/api/runs/{run['id']}").json()

    config = RunConfig(**reopened["run_config"])
    assert config.knowledge_base_version_id == upload["project"]["active_version_id"]
    assert config.llm_model == DEFAULT_OLLAMA_LLM_MODEL
    assert config.vlm_model == DEFAULT_OLLAMA_VLM_MODEL
    assert config.embedding_model == "bge-m3"
    assert all(result["prompt_versions"] for result in reopened["results"])
    traditional = next(result for result in reopened["results"] if result["pipeline_type"] == "traditional")
    assert traditional["prompt_versions"]["answer"] == "traditional_answer_v1"


def test_epic12_prompt_policy_catches_uncited_factual_answers():
    warnings = validate_answer_contract("Based on the uploaded evidence:\n- Authentication uses API Gateway.", [])
    assert "answer_claims_without_citations" in warnings
    assert validate_answer_contract("I could not find enough evidence in the uploaded sources to answer this question.", []) == []


def test_epic13_privacy_defaults_disable_cloud_telemetry_and_paid_requirements():
    privacy = client.get("/api/privacy")
    dependencies = client.get("/api/settings/dependencies")
    models = client.get("/api/settings/models")

    assert privacy.status_code == 200
    assert privacy.json()["local_first"] is True
    assert privacy.json()["telemetry_enabled"] is False
    assert privacy.json()["cloud_providers_enabled"] is False
    assert privacy.json()["paid_api_required"] is False
    assert dependencies.json()["telemetry"]["enabled"] is False
    assert dependencies.json()["cloud_providers"]["enabled"] is False
    assert "Local model outputs can be incorrect" in models.json()["local_model_warning"]["warning"]


def test_epic13_project_delete_removes_artifacts_and_database_rows():
    project = client.post("/api/projects", json={"name": "Delete Lifecycle"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway.", "text/plain")},
    ).json()
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What uses API Gateway?"},
    ).json()
    project_path = os.path.join(os.environ["DATA_DIR"], "projects", project["id"])
    assert os.path.exists(project_path)

    deleted = client.delete(f"/api/projects/{project['id']}")

    assert deleted.status_code == 200
    assert not os.path.exists(project_path)
    with connect() as conn:
        tables = ["sources", "content_blocks", "chunks", "vector_records", "runs", "pipeline_results", "entities", "relationships", "graph_states"]
        for table in tables:
            row = conn.execute(f"SELECT COUNT(*) AS count FROM {table} WHERE project_id = ?", (project["id"],)).fetchone() if table not in {"pipeline_results", "graph_states"} else None
            if row is not None:
                assert row["count"] == 0
        assert conn.execute("SELECT COUNT(*) AS count FROM pipeline_results WHERE run_id = ?", (run["id"],)).fetchone()["count"] == 0
        assert conn.execute("SELECT COUNT(*) AS count FROM graph_states WHERE run_id = ?", (run["id"],)).fetchone()["count"] == 0
    assert client.get(f"/api/sources/{upload['sources'][0]['id']}").status_code == 404


def test_epic13_exports_are_scrubbed_and_include_privacy_flags():
    project = client.post("/api/projects", json={"name": "Private Export"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What uses API Gateway?"},
    ).json()

    exported = client.get(f"/api/runs/{run['id']}/export.json").json()
    text = client.get(f"/api/runs/{run['id']}/export.md").text

    assert exported["privacy"]["cloud_provider_used"] is False
    assert exported["privacy"]["telemetry_enabled"] is False
    assert "/tmp/ragbench-test-data" not in str(exported)
    assert "api_key" not in str(exported).lower()
    assert "/tmp/ragbench-test-data" not in text


def test_epic14_sample_dataset_loads_sources_and_questions():
    info = client.get("/api/sample-dataset")
    loaded = client.post("/api/sample-dataset/load")

    assert info.status_code == 200
    assert len(info.json()["questions"]) == 6
    assert {
        "01_northstar_incident_brief.md",
        "02_service_tickets.csv",
        "03_service_review_notes.md",
    }.issubset(set(info.json()["sources"]))
    assert loaded.status_code == 200
    project = loaded.json()["project"]
    assert project["source_count"] == 3
    assert project["chunk_count"] >= 5
    assert project["active_version"]["change_type"] == "sample_dataset"


def test_epic14_acceptance_suite_runs_and_records_timings():
    result = client.post("/api/acceptance/run")

    assert result.status_code == 200
    body = result.json()
    assert body["passed"] is True
    assert body["elapsed_seconds"] >= 0
    assert {item["id"] for item in body["results"]} == {
        "root_cause_and_risk",
        "firmware_customer_impact",
        "rollback_decision",
        "site_comparison",
        "owners_and_actions",
        "false_positive_case",
    }
    assert all(item["elapsed_seconds"] >= 0 for item in body["results"])
    assert all(item["expected_hits"] for item in body["results"])


def test_epic15_guardrails_disable_excluded_mvp_scope_and_track_deferred_scope():
    response = client.get("/api/guardrails")

    assert response.status_code == 200
    body = response.json()
    assert body["credentials_required"] is False
    assert body["paid_or_cloud_default"] is False
    assert body["local_first_required"] is True
    assert set(body["mvp_exclusions"]) == set(MVP_EXCLUSIONS)
    assert all(item["enabled"] is False for item in body["mvp_exclusions"].values())
    assert set(REQUIRED_DEFERRED_SCOPE).issubset(set(body["required_deferred_scope"]))
    assert all(item["tracked"] is True for item in body["required_deferred_scope"].values())


def test_epic15_forbidden_defaults_are_detected_and_credentials_are_not_required():
    assert assert_no_forbidden_defaults({"paid_api_required": True, "credentials_required": True}) == [
        "paid_api_required",
        "credentials_required",
    ]
    privacy = client.get("/api/privacy").json()
    guardrails = client.get("/api/guardrails").json()
    providers = client.get("/api/settings/providers").json()

    assert privacy["paid_api_required"] is False
    assert guardrails["credentials_required"] is False
    assert "api_key" not in str(providers).lower()


def test_epic16_ingestion_job_contract_and_events():
    project = client.post("/api/projects", json={"name": "Ingestion Jobs"}).json()

    response = client.post(f"/api/projects/{project['id']}/ingestion-jobs", json={"source_ids": ["source-1"]})

    assert response.status_code == 200
    job = response.json()
    assert job["status"] == "succeeded"
    assert {job["status"], "queued", "running", "failed", "cancelled", "partial_success"}.issubset(set(JOB_STATES))
    events = client.get(f"/api/jobs/{job['id']}/events").json()
    assert [event["event_type"] for event in events] == list(INGESTION_STEPS)
    assert client.get(f"/api/jobs/{job['id']}").json()["id"] == job["id"]


def test_epic16_benchmark_job_steps_sse_websocket_and_cancel():
    project = client.post("/api/projects", json={"name": "Benchmark Jobs"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()

    events = client.get(f"/api/runs/{run['id']}/events").json()
    assert [event["event_type"] for event in events] == list(BENCHMARK_STEPS)
    sse = client.get(f"/api/runs/{run['id']}/events.sse")
    assert sse.status_code == 200
    assert "data:" in sse.text
    with client.websocket_connect(f"/api/runs/{run['id']}/events.ws") as websocket:
        first_event = websocket.receive_json()
    assert first_event["event_type"] == "create_run"

    cancelled = client.post(f"/api/runs/{run['id']}/cancel")
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"


def test_epic16_job_sse_and_websocket_parity():
    project = client.post("/api/projects", json={"name": "Job Streams"}).json()
    job = client.post(f"/api/projects/{project['id']}/ingestion-jobs", json={}).json()

    sse = client.get(f"/api/jobs/{job['id']}/events.sse")
    assert sse.status_code == 200
    assert "validate_file" in sse.text
    with client.websocket_connect(f"/api/jobs/{job['id']}/events.ws") as websocket:
        first_event = websocket.receive_json()
    assert first_event["event_type"] == "validate_file"


def test_epic16_queue_runner_contracts_and_model_settings():
    runners = [LocalJobRunner(), RedisRQJobRunner(), CeleryJobRunner()]
    assert {runner.contract_name for runner in runners} == {"local_background_tasks", "redis_rq", "celery"}
    assert all(callable(runner.enqueue) and callable(runner.cancel) and callable(runner.status) for runner in runners)

    patched = client.patch("/api/settings/models", json={"llm_model": "qwen3:14b", "embedding_model": "bge-m3"})
    checked = client.post("/api/settings/model-check", json={"provider": "ollama", "model_type": "llm", "model": "qwen3:14b"})

    assert patched.status_code == 200
    assert patched.json()["overrides"]["llm_model"] == "qwen3:14b"
    assert checked.status_code == 200
    assert checked.json()["container_required"] is False
    assert checked.json()["host_ollama_required"] is True


def test_model_settings_report_and_pull_missing_host_ollama_models(monkeypatch):
    main_module.MODEL_SETTINGS_OVERRIDES.clear()
    project = client.post("/api/projects", json={"name": "Model Pull"}).json()
    installed = {"qwen3-vl:4b", "bge-m3"}
    pulled = []

    def fake_installed_models(_base_url, timeout=10.0):
        return set(installed)

    def fake_pull_model(_base_url, model):
        pulled.append(model)
        installed.add(model)

    monkeypatch.setattr(main_module, "ollama_installed_models", fake_installed_models)
    monkeypatch.setattr(main_module, "ollama_pull_model", fake_pull_model)
    monkeypatch.setattr(main_module, "wait_for_ollama", lambda _base_url, attempts=10: None)

    status = client.get("/api/settings/models").json()["ollama_model_management"]
    assert status["missing_models"] == ["qwen3:8b"]
    assert status["pull_commands"] == ["ollama pull qwen3:8b"]

    pull = client.post(
        "/api/settings/models/pull",
        json={"project_id": project["id"], "models": ["qwen3:8b"]},
    ).json()

    assert pull["job_status"] == "succeeded"
    assert pulled == ["qwen3:8b"]
    assert client.get("/api/settings/models/ollama").json()["missing_models"] == []


def test_database_dashboards_expose_local_query_tools():
    dashboards = client.get("/api/settings/database-dashboards").json()["dashboards"]
    by_id = {item["id"]: item for item in dashboards}

    assert by_id["qdrant"]["url"] == "http://localhost:6333/dashboard"
    assert by_id["chroma"]["url"] == "http://localhost:8001/docs"
    assert by_id["sqlite"]["url"] == "http://localhost:8000/api/database/sqlite-dashboard"
    assert by_id["redis"]["url"] == "http://localhost:8083"
    assert all(item["queryable"] is True for item in dashboards)

    observability = client.get("/api/settings/observability").json()
    observability_by_id = {item["id"]: item for item in observability["dashboards"]}
    assert observability_by_id["jaeger"]["url"] == "http://localhost:16686"
    assert observability_by_id["phoenix"]["url"] == "http://localhost:6006"
    assert observability_by_id["langsmith"]["external"] is True
    assert observability["langsmith"]["project"]

    query = client.get("/api/database/sqlite-query", params={"sql": "SELECT name FROM sqlite_master WHERE type = 'table' ORDER BY name"})
    blocked = client.get("/api/database/sqlite-query", params={"sql": "DELETE FROM projects"})

    assert query.status_code == 200
    assert query.json()["row_count"] >= 1
    assert blocked.status_code == 400


def test_epic16_partial_pipeline_failure_keeps_successful_results_visible(monkeypatch):
    project = client.post("/api/projects", json={"name": "Partial Run"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    original = main_module.run_pipeline

    def flaky_pipeline(project_id, question, pipeline_type, selected_source_ids=None):
        if pipeline_type == "agentic":
            raise RuntimeError("agentic host Ollama unavailable")
        return original(project_id, question, pipeline_type, selected_source_ids=selected_source_ids)

    monkeypatch.setattr(main_module, "run_pipeline", flaky_pipeline)
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    )

    assert run.status_code == 200
    body = run.json()
    assert len(body["results"]) == 3
    failed = next(result for result in body["results"] if result["pipeline_type"] == "agentic")
    traditional = next(result for result in body["results"] if result["pipeline_type"] == "traditional")
    assert "pipeline_failed" in failed["warnings"]
    assert traditional["citations"]
    assert client.get(f"/api/runs/{body['id']}/events").json()[-1]["status"] == "partial_success"


def test_epic17_storage_layout_manifest_migrations_and_trace_replay():
    project = client.post("/api/projects", json={"name": "Storage Layout"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    ).json()
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()

    layout = client.get("/api/storage/layout").json()["layout"]
    migrations = client.get("/api/settings/migrations").json()
    trace_text = client.get(f"/api/runs/{run['id']}/trace.jsonl").text
    trace_events = client.get(f"/api/runs/{run['id']}/trace-events").json()

    assert layout["vector_store"] == ARTIFACT_LAYOUT["vector_store"]
    assert migrations["current_version"] == LATEST_SCHEMA_VERSION
    source = client.get(f"/api/sources/{upload['sources'][0]['id']}").json()
    assert source["local_path"].startswith(f"projects/{project['id']}/sources/original/")
    assert run["artifact_paths"]["traditional"].startswith(f"projects/{project['id']}/runs/{run['id']}/")
    assert os.path.exists(os.path.join(os.environ["DATA_DIR"], "projects", project["id"], "project.json"))
    assert parse_trace_jsonl(trace_text)
    assert trace_events["event_count"] == len(parse_trace_jsonl(trace_text))
    reopened = client.get(f"/api/runs/{run['id']}").json()
    assert reopened["run_config"]["knowledge_base_version_id"] == upload["project"]["active_version_id"]


def test_epic17_migration_status_on_fresh_database():
    with connect() as conn:
        status = migration_status(conn)
    assert status["current_version"] == LATEST_SCHEMA_VERSION
    assert status["pending_versions"] == []


def test_epic18_reliability_defaults_disk_warning_and_performance_smoke(monkeypatch):
    monkeypatch.setenv("RAGBENCH_DISK_WARNING_BYTES", str(10**18))
    reliability = client.get("/api/settings/reliability").json()
    project = client.post("/api/projects", json={"name": "Perf Smoke"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    )
    model_check = client.post("/api/settings/model-check", json={"provider": "ollama", "model_type": "llm", "model": "qwen3:14b"}).json()

    assert reliability["defaults"]["top_k_chunks"] == DEFAULT_TOP_K_CHUNKS == 8
    assert reliability["defaults"]["reranked_chunks"] == DEFAULT_RERANKED_CHUNKS == 5
    assert reliability["disk"]["warning"] is True
    assert {"processing", "completed", "failed", "waiting_for_local_models", "partial_success"}.issubset(set(reliability["processing_states"]))
    assert upload.status_code == 200
    assert run.status_code == 200
    assert run.json()["results"][0]["metrics"]["top_k_chunks"] == 8
    assert run.json()["results"][0]["metrics"]["reranked_chunks"] == 5
    assert model_check["progress"]["status"] == "waiting_for_local_models"


def test_resource_profile_preflight_selects_safe_model_flow(monkeypatch):
    monkeypatch.setenv("RAGBENCH_MODEL_PROFILE", "auto")
    lite_resources = {"free_disk_bytes": 45 * GIB, "memory_bytes": 16 * GIB, "free_disk_gib": 45, "memory_gib": 16}
    lite = preflight_report(lite_resources)
    assert lite["safe_to_pull"] is True
    assert lite["selected_profile"] == "lite"
    assert lite["models_to_pull"] == ["qwen3:8b", "qwen3-vl:4b", "bge-m3"]

    small_docker_resources = {"free_disk_bytes": 10 * GIB, "memory_bytes": 8 * GIB, "free_disk_gib": 10, "memory_gib": 8}
    host_managed = preflight_report(small_docker_resources)
    assert host_managed["safe_to_pull"] is True
    assert host_managed["runtime"] == "host_ollama"
    assert "ollama pull qwen3:8b" in host_managed["pull_commands"]

    endpoint = client.get("/api/settings/resource-profile")
    assert endpoint.status_code == 200
    assert "profiles" in endpoint.json()


def test_epic19_advanced_rag_variants_are_activated_and_capped():
    project = client.post("/api/projects", json={"name": "Advanced Variants"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("facts.txt", b"Authentication: API Gateway | Token Store. Deployment latency risk uses monitoring.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does Authentication use?"},
    ).json()
    traditional = next(result for result in run["results"] if result["pipeline_type"] == "traditional")
    variants = traditional["metrics"]["advanced_rag_variants"]

    for name in [
        "advanced_rag",
        "hybrid_search_rag",
        "fusion_rag",
        "self_rag",
        "hierarchical_rag",
        "structured_rag",
        "long_context_rag",
        "adaptive_rag",
    ]:
        assert name in variants
    assert len(variants["fusion_rag"]["query_variants"]) <= 3
    assert variants["long_context_rag"]["context"]["citations"]
    assert "fusion_rag" in traditional["techniques"]
    assert any(step["step"] == "advanced_rag_variants" for step in traditional["trace"])


def test_epic19_structured_and_long_context_preserve_citations():
    project = client.post("/api/projects", json={"name": "Structured Long Context"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("table.md", b"| Component | Owner |\n| API Gateway | Platform |\n| Token Store | Security |", "text/markdown")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "Who owns API Gateway?"},
    ).json()
    traditional = next(result for result in run["results"] if result["pipeline_type"] == "traditional")
    variants = traditional["metrics"]["advanced_rag_variants"]

    assert variants["structured_rag"]["activated"] is True
    assert variants["structured_rag"]["facts_count"] >= 1
    assert variants["long_context_rag"]["context"]["token_estimate"] <= variants["long_context_rag"]["context"]["max_tokens"]
    assert all(citation["label"] in variants["long_context_rag"]["context"]["context"] for citation in traditional["citations"])


def test_epic20_vector_metadata_graph_and_media_adapter_contract_parity():
    assert equivalent_vector_results(ChromaVectorStoreAdapter(), QdrantVectorStoreAdapter(), "Authentication API Gateway")

    sqlite_repo = SQLiteMetadataRepository()
    sqlite_repo.create_project("p1", "Adapter Project")
    assert sqlite_repo.get_project("p1")["name"] == "Adapter Project"
    assert sqlite_repo.list_projects()

    assert equivalent_graph_results(SQLiteGraphStoreAdapter(), NetworkXGraphStoreAdapter())

    ocr = normalize_ocr_output("easyocr", "API Gateway", 0.8, [{"box": [0, 0, 1, 1]}])
    transcript = normalize_transcript_output("faster-whisper", [{"text": "API Gateway", "start": 0, "end": 1, "confidence": 0.9}])
    assert set(ocr) == {"provider", "text", "confidence", "regions", "warnings"}
    assert transcript["segments"][0]["timestamp_start"] == 0.0
    assert transcript["segments"][0]["timestamp_end"] == 1.0


def test_epic20_settings_expose_adapter_choices():
    adapters = client.get("/api/settings/adapters")

    assert adapters.status_code == 200
    body = adapters.json()
    assert body["vector_stores"]["qdrant"]["container"] == "qdrant"
    assert body["metadata_stores"]["sqlite"]["required"] is True
    assert body["metadata_stores"]["postgres"]["required"] is False
    assert body["graph_stores"]["networkx"]["provider"] == "networkx"
    assert body["runtimes"]["ollama"]["runtime"] == "host"
    assert body["ocr"]["easyocr"]["provider"] == "easyocr"
    assert body["transcription"]["faster_whisper"]["provider"] == "faster-whisper"


def test_epic21_follow_up_keeps_uploaded_sources_authoritative():
    project = client.post("/api/projects", json={"name": "Follow Up Authority"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store. Generated answers are not source evidence.", "text/plain")},
    )
    parent = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    ).json()
    follow_up = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "Which source supports that?", "parent_run_id": parent["id"], "run_mode": "follow_up"},
    ).json()

    assert follow_up["parent_run_id"] == parent["id"]
    assert follow_up["user_feedback"] is None
    for result in follow_up["results"]:
        for citation in result["citations"]:
            assert citation["source_id"] != parent["id"]
            assert citation["label"].startswith("auth.txt#")


def test_epic21_conversation_mode_does_not_index_generated_answers():
    project = client.post("/api/projects", json={"name": "Conversation"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    before_chunks = client.get(f"/api/projects/{project['id']}").json()["chunk_count"]
    session = client.post(
        f"/api/projects/{project['id']}/conversation-sessions",
        json={"title": "Auth conversation"},
    ).json()
    updated = client.post(
        f"/api/conversation-sessions/{session['id']}/messages",
        json={"question": "What does authentication use?"},
    ).json()
    after_project = client.get(f"/api/projects/{project['id']}").json()

    assert updated["summary"].startswith("Follow-up context")
    assert len(updated["messages"]) == 2
    assistant = next(message for message in updated["messages"] if message["role"] == "assistant")
    assert assistant["generated_answer_indexed"] is False
    assert assistant["source_citations"]
    assert after_project["chunk_count"] == before_chunks


def test_epic21_app_modes_desktop_and_browser_only_feature_matrix():
    response = client.get("/api/app-modes")

    assert response.status_code == 200
    modes = response.json()
    assert modes["conversation_mode"]["generated_answers_indexed"] is False
    assert modes["desktop_app"]["local_first_storage"] is True
    assert modes["desktop_app"]["cloud_required"] is False
    assert modes["browser_only_local_mode"]["feature_matrix"]["text_markdown_parsing"] == "supported_in_browser"
    assert modes["browser_only_local_mode"]["feature_matrix"]["local_llm_vlm_inference"] == "requires_host_ollama"


def test_epic22_paid_cloud_extensions_are_disabled_by_default_and_do_not_block_workflows():
    extensions = client.get("/api/settings/provider-extensions").json()
    project = client.post("/api/projects", json={"name": "Local Without Paid"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(
        f"/api/projects/{project['id']}/runs",
        json={"question": "What does authentication use?"},
    )
    export = client.get(f"/api/runs/{run.json()['id']}/export.json")

    assert assert_local_first_provider_policy(extensions) == []
    assert extensions["defaults"]["paid_cloud_selected"] is False
    assert extensions["defaults"]["credentials_required_for_local_workflows"] is False
    assert all(provider["enabled"] is False for provider in extensions["paid_cloud_extensions"].values())
    assert upload.status_code == 200
    assert run.status_code == 200
    assert export.status_code == 200


def test_epic22_paid_cloud_stub_contract_never_calls_network_without_explicit_enablement():
    calls = []

    def network_call(payload):
        calls.append(payload)
        return {"ok": True}

    stub = PaidCloudProviderStub("openai", "llm", enabled=False, network_call=network_call)
    disabled = stub.execute({"uploaded_text": "secret evidence"}, explicit_user_enabled=False)
    missing_credentials = PaidCloudProviderStub("openai", "llm", enabled=True, network_call=network_call).execute(
        {"uploaded_text": "secret evidence"},
        explicit_user_enabled=True,
    )
    configured = PaidCloudProviderStub("openai", "llm", enabled=True, network_call=network_call).execute(
        {"prompt": "hello"},
        explicit_user_enabled=True,
        credentials="token",
    )

    assert disabled["network_called"] is False
    assert missing_credentials["network_called"] is False
    assert configured["network_called"] is True
    assert calls == [{"prompt": "hello"}]
    settings = extension_settings()
    assert set(settings["extension_types"]) == {"llm", "vlm", "ocr", "transcription", "vector_database", "object_storage"}


def test_epic22_settings_separate_local_required_from_paid_cloud_extensions():
    adapters = client.get("/api/settings/adapters").json()
    extensions = client.get("/api/settings/provider-extensions").json()

    assert "paid_cloud_extensions" in adapters
    assert "local_required_providers" in extensions
    assert "paid_cloud_extensions" in extensions
    assert extensions["local_required_providers"]["object_storage"] == "local_project_artifacts"
    assert extensions["defaults"]["send_uploaded_data_without_explicit_configuration"] is False


def test_epic23_provider_health_reports_real_mode_and_host_ollama_remediation(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    body = client.get("/api/settings/provider-health").json()

    assert body["real_mode"] is True
    assert body["ollama"]["runtime"] == "host"
    assert "Install Ollama" in body["ollama"]["remediation"]
    assert "storage" in body
    assert "queues" in body
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "deterministic")


def test_epic24_deterministic_mode_is_explicit_and_model_check_has_host_ollama_setup():
    models = client.get("/api/settings/models").json()
    checked = client.post("/api/settings/model-check", json={"provider": "ollama", "model_type": "embedding", "model": "bge-m3"}).json()

    assert models["embeddings"]["provider"] == "deterministic_lexical"
    assert checked["container_required"] is False
    assert checked["host_ollama_required"] is True
    assert checked["setup"].startswith("ollama pull")


def test_epic25_vector_store_adapters_support_delete_lifecycle():
    records = [
        {"id": "v1", "source_id": "s1", "chunk_id": "c1", "text": "Authentication uses API Gateway.", "metadata": {"source_id": "s1", "chunk_id": "c1"}},
        {"id": "v2", "source_id": "s2", "chunk_id": "c2", "text": "Deployment latency budget.", "metadata": {"source_id": "s2", "chunk_id": "c2"}},
    ]
    store = ChromaVectorStoreAdapter(project_id="fixture")
    store.upsert(records)
    assert store.query("Authentication API", top_k=1)
    store.delete_source("s1")
    assert all(hit.get("source_id") != "s1" for hit in store.query("Authentication API", top_k=5))
    store.delete_project()


def test_epic26_real_mode_rejects_invalid_image_without_placeholder(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    project = client.post("/api/projects", json={"name": "Real Media Failure"}).json()
    response = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("bad.png", b"not-a-real-image", "image/png")},
    ).json()

    assert response["sources"][0]["status"] == "failed"
    assert "Deterministic OCR placeholder" not in str(response)
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "deterministic")


def test_epic27_graph_state_uses_actual_trace_nodes_not_prefilled_list():
    project = client.post("/api/projects", json={"name": "Actual Graph Nodes"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()
    states = client.get(f"/api/runs/{run['id']}/graph-states").json()
    final_states = {}
    for state in states:
        final_states[state["pipeline_type"]] = state
        assert state["state"]["completed_nodes"]
        assert state["state"]["metrics"]["graph_runtime"] == "langgraph"
    for state in final_states.values():
        assert state["state"]["metrics"]["graph_nodes_executed"] >= len(GRAPH_NODES[state["pipeline_type"]])


def test_epic37_graph_persists_node_state_snapshots():
    project = client.post("/api/projects", json={"name": "Graph Snapshots"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()
    states = client.get(f"/api/runs/{run['id']}/graph-states").json()
    traditional_states = [state for state in states if state["pipeline_type"] == "traditional"]

    assert len(traditional_states) > 1
    assert traditional_states[0]["state"]["metrics"]["graph_snapshot_index"] == 1
    assert traditional_states[-1]["state"]["metrics"]["graph_nodes_executed"] >= len(GRAPH_NODES["traditional"])
    assert any(step.get("graph_node_scope") == "node_state_transition" for step in run["results"][0]["trace"])


def test_epic28_queued_benchmark_is_processed_by_worker(monkeypatch):
    project = client.post("/api/projects", json={"name": "Queued Benchmark"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    monkeypatch.setenv("RAGBENCH_QUEUE_MODE", "local")
    queued = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()
    assert queued["job_status"] == "queued"

    job = fetch_next_queued_job()
    assert job["job_type"] == "benchmark"
    process_job(job)
    completed = client.get(f"/api/runs/{queued['id']}").json()
    assert completed["job_status"] == "succeeded"
    assert len(completed["results"]) == 3
    monkeypatch.setenv("RAGBENCH_QUEUE_MODE", "inline")


def test_epic28_source_upload_is_staged_and_processed_by_worker(monkeypatch):
    monkeypatch.setenv("RAGBENCH_QUEUE_MODE", "local")
    project = client.post("/api/projects", json={"name": "Queued Sources"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    ).json()
    assert upload["job_status"] == "queued"
    assert upload["sources"][0]["status"] == "queued"

    job = fetch_next_queued_job()
    assert job["job_type"] == "source_upload"
    process_job(job)
    completed = client.get(f"/api/jobs/{job['id']}").json()
    project_after = client.get(f"/api/projects/{project['id']}").json()
    assert completed["status"] == "succeeded"
    assert project_after["chunk_count"] == 1
    monkeypatch.setenv("RAGBENCH_QUEUE_MODE", "inline")


def test_epic29_schema_version_and_repository_status_are_real():
    migrations = client.get("/api/settings/migrations").json()
    repositories = client.get("/api/settings/repositories").json()

    assert migrations["latest_schema_version"] == 34
    assert migrations["current_version"] == 34
    assert repositories["sqlite"]["schema_managed"] is True
    assert repositories["postgres"]["required"] is False


def test_epic30_graph_endpoint_exposes_typed_entities_relationships_and_provenance():
    project = client.post("/api/projects", json={"name": "Graph Endpoint"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store for sessions.", "text/plain")},
    )
    graph = client.get(f"/api/projects/{project['id']}/graph").json()

    assert graph["nodes"]
    assert graph["edges"]
    assert "aliases" in graph["nodes"][0]
    assert graph["edges"][0]["evidence_chunk_id"]
    assert graph["edges"][0]["metadata"]["provenance"] == "chunk"


def test_epic31_advanced_rag_trace_proves_work_performed():
    project = client.post("/api/projects", json={"name": "Advanced Work"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("table.md", b"Component | Owner\nAPI Gateway | Platform\nToken Store | Security", "text/markdown")},
    )
    result = run_pipeline(project["id"], "Who owns API Gateway?", "traditional")
    steps = {step["step"] for step in result["trace"]}

    assert "corrective_rag_branch" in steps
    assert "structured_rag_lookup" in steps
    assert "self_rag_critic" in steps
    assert result["metrics"]["structured_facts_count"] >= 1


def test_epic32_citation_resolver_returns_source_block_and_artifact_metadata():
    project = client.post("/api/projects", json={"name": "Citation Resolve"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()
    citation = run["results"][0]["citations"][0]
    resolved = client.get(f"/api/citations/resolve?chunk_id={citation['chunk_id']}&run_id={run['id']}").json()
    all_resolved = client.get(f"/api/runs/{run['id']}/citations").json()

    assert resolved["source_id"] == citation["source_id"]
    assert resolved["text"]
    assert resolved["evidence_kind"] in {"source_text", "graph_relationship", "ocr_visual", "transcript", "structured_table"}
    assert all_resolved


def test_epic33_app_modes_have_real_browser_and_desktop_entrypoints():
    modes = client.get("/api/app-modes").json()

    assert modes["desktop_app"]["smoke_command"] == "cd apps/desktop && npm run smoke"
    assert modes["browser_only_local_mode"]["entrypoint"] == "apps/web/browser-only.html"
    assert modes["browser_only_local_mode"]["feature_matrix"]["image_audio_video_processing"] == "requires_backend_container"


def test_epic34_no_gap_audit_records_final_acceptance_report():
    report = client.post("/api/acceptance/no-gap-audit").json()

    assert report["status"] in {"passed", "failed"}
    assert set(report["required_epics"]) == set(range(23, 51))
    assert report["checks"]["browser_only_exists"] is True
    assert report["checks"]["desktop_wrapper_exists"] is True


class FakeOCR:
    provider = "fake_ocr"

    def extract_text(self, image_path):
        return ProviderResult(provider=self.provider, model="fake-ocr-v1", text=f"OCR text from {image_path.name}", metadata={"confidence": 0.91})


class FakeVLM:
    provider = "fake_vlm"
    model = "fake-vlm-v1"

    def inspect_image(self, source_id, frame_or_image_id, image_path, question):
        return ProviderVisualObservation(
            source_id=source_id,
            frame_or_image_id=frame_or_image_id,
            question=question,
            visual_answer=f"VLM caption for {image_path.name}: API Gateway points to Token Store.",
            confidence=0.87,
            provider=self.provider,
            model=self.model,
        )


class FakeRegistry:
    def ocr(self):
        return FakeOCR()

    def vlm(self):
        return FakeVLM()


class SequenceVLM:
    provider = "fake_vlm"
    model = "fake-vlm-v1"

    def __init__(self, answers):
        self.answers = list(answers)
        self.questions = []

    def inspect_image(self, source_id, frame_or_image_id, image_path, question):
        self.questions.append(question)
        answer, confidence = self.answers.pop(0)
        return ProviderVisualObservation(
            source_id=source_id,
            frame_or_image_id=frame_or_image_id,
            question=question,
            visual_answer=answer,
            confidence=confidence,
            provider=self.provider,
            model=self.model,
        )


class VLMOnlyRegistry:
    def __init__(self, vlm):
        self._vlm = vlm

    def vlm(self):
        return self._vlm


class FakeLLM:
    provider = "fake_llm"
    model = "fake-llm-v1"

    def __init__(self, responses):
        self.responses = list(responses)
        self.prompts = []

    def generate_structured(self, prompt, schema):
        self.prompts.append(prompt)
        return self.responses.pop(0)


class LLMOnlyRegistry:
    def __init__(self, llm):
        self._llm = llm

    def llm(self):
        return self._llm


def insert_visual_source(project_id: str, source_id: str, ocr_text: str = "Login Service connects to Token Store.") -> None:
    image_path = data_dir() / "projects" / project_id / "sources" / "original" / f"{source_id}.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(b"fake image")
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO sources (id, project_id, filename, source_type, mime_type, local_path, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (source_id, project_id, f"{source_id}.png", "image", "image/png", str(image_path.relative_to(data_dir())), "processed", "2026-04-26T00:00:00Z"),
        )
        conn.execute(
            """
            INSERT INTO content_blocks (id, project_id, source_id, block_type, text, confidence, metadata_json)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (f"{source_id}-ocr", project_id, source_id, "ocr", ocr_text, 0.9, "{}"),
        )


def test_epic35_real_image_ingestion_creates_vlm_caption_block(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    monkeypatch.setattr(media_processing, "provider_registry", lambda: FakeRegistry())
    project_root = data_dir() / "projects" / "epic35-image"
    image_path = project_root / "sources" / "original" / "diagram.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(b"fake image bytes inspected by fake VLM")

    blocks, warnings = media_processing.image_blocks(image_path, "diagram.png", project_root, "source-image")

    assert not warnings
    block_types = [block["block_type"] for block in blocks]
    assert "ocr" in block_types
    assert "caption" in block_types
    caption = next(block for block in blocks if block["block_type"] == "caption")
    assert caption["metadata"]["provider"] == "fake_vlm"
    assert caption["metadata"]["model"] == "fake-vlm-v1"
    assert caption["metadata"]["visual_caption"] is True
    assert "Deterministic caption placeholder" not in caption["text"]


def test_host_media_runtime_path_uses_host_ocr_without_container_tesseract(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    monkeypatch.setenv("RAGBENCH_MEDIA_RUNTIME", "host")
    monkeypatch.setattr(media_processing, "provider_registry", lambda: FakeRegistry())

    class FakeHostMediaClient:
        def ocr(self, path):
            return {"provider": "host-tesseract", "model": "tesseract-cli", "text": "Host OCR text", "confidence": 0.91}

    monkeypatch.setattr(media_processing, "HostMediaClient", lambda: FakeHostMediaClient())
    project_root = data_dir() / "projects" / "host-media-image"
    image_path = project_root / "sources" / "original" / "diagram.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(b"fake image bytes")

    blocks, warnings = media_processing.image_blocks(image_path, "diagram.png", project_root, "source-image")

    assert not warnings
    ocr = next(block for block in blocks if block["block_type"] == "ocr")
    assert ocr["text"] == "Host OCR text"
    assert ocr["metadata"]["provider"] == "host-tesseract"


def test_epic35_real_video_ingestion_creates_frame_vlm_caption_blocks(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    monkeypatch.setattr(media_processing, "provider_registry", lambda: FakeRegistry())
    monkeypatch.setattr(media_processing, "transcribe_audio", lambda path, project_root, source_id: ([{"block_type": "transcript", "text": "checkout issue mentioned", "timestamp_start": 0.0, "timestamp_end": 1.0, "confidence": 0.8, "metadata": {"provider": "fake"}}], []))

    project_root = data_dir() / "projects" / "epic35-video"
    video_path = project_root / "sources" / "original" / "demo.mp4"
    video_path.parent.mkdir(parents=True, exist_ok=True)
    video_path.write_bytes(b"fake video bytes")
    audio_path = project_root / "sources" / "derived" / "audio" / "source-video.wav"
    frame_dir = project_root / "sources" / "derived" / "frames" / "source-video"
    audio_path.parent.mkdir(parents=True, exist_ok=True)
    frame_dir.mkdir(parents=True, exist_ok=True)
    audio_path.write_bytes(b"audio")
    frames = []
    for index in range(2):
        frame = frame_dir / f"frame_{index + 1:04d}.jpg"
        frame.write_bytes(b"frame")
        frames.append(frame)
    monkeypatch.setattr(media_processing, "extract_audio", lambda path, project_root, source_id: audio_path)
    monkeypatch.setattr(media_processing, "extract_frames", lambda path, project_root, source_id: frames)
    monkeypatch.setattr(media_processing, "TesseractOCRAdapter", lambda: FakeOCR())

    blocks, warnings = media_processing.video_blocks(video_path, "demo.mp4", project_root, "source-video")

    assert not warnings
    captions = [block for block in blocks if block["block_type"] == "caption"]
    assert len(captions) == 2
    assert all(caption["frame_path"] for caption in captions)
    assert all(caption["metadata"]["provider"] == "fake_vlm" for caption in captions)
    assert all("Deterministic frame caption placeholder" not in caption["text"] for caption in captions)


def test_epic36_weak_vlm_answer_triggers_one_refined_follow_up(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    project = client.post("/api/projects", json={"name": "VLM Retry"}).json()
    insert_visual_source(project["id"], "visual-retry")
    vlm = SequenceVLM(
        [
            ("The image shows a system diagram.", 0.3),
            ("The authentication diagram shows Login Service issuing tokens through Token Store.", 0.86),
        ]
    )
    monkeypatch.setattr(agentic_module, "provider_registry", lambda: VLMOnlyRegistry(vlm))
    state = AgentState(project_id=project["id"], question="What authentication flow is shown?")
    observation = ToolRegistry(state).call(
        "inspect_visual_source",
        source_id="visual-retry",
        frame_or_image_id=None,
        question_for_vlm=state.question,
    )

    assert observation["follow_up_performed"] is True
    assert state.vlm_calls == 2
    assert len(observation["observations"]) == 2
    assert "Login Service" in observation["text"]
    assert any(step["step"] == "vlm_follow_up" for step in state.trace)
    assert vlm.questions[1] != vlm.questions[0]


def test_epic36_strong_vlm_answer_does_not_retry(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    project = client.post("/api/projects", json={"name": "VLM No Retry"}).json()
    insert_visual_source(project["id"], "visual-strong")
    vlm = SequenceVLM(
        [
            ("The authentication flow shows Login Service connecting to Token Store for issued tokens.", 0.91),
            ("This answer should not be used.", 0.1),
        ]
    )
    monkeypatch.setattr(agentic_module, "provider_registry", lambda: VLMOnlyRegistry(vlm))
    state = AgentState(project_id=project["id"], question="What authentication flow is shown?")
    observation = ToolRegistry(state).call(
        "inspect_visual_source",
        source_id="visual-strong",
        frame_or_image_id=None,
        question_for_vlm=state.question,
    )

    assert observation["follow_up_performed"] is False
    assert state.vlm_calls == 1
    assert len(observation["observations"]) == 1
    assert len(vlm.questions) == 1


def test_epic36_ocr_vlm_conflict_is_reported(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    project = client.post("/api/projects", json={"name": "VLM Conflict"}).json()
    insert_visual_source(project["id"], "visual-conflict", ocr_text="Login Service connects to Token Store.")
    vlm = SequenceVLM([("The authentication issue is a fruit basket checkout screen.", 0.9)])
    monkeypatch.setattr(agentic_module, "provider_registry", lambda: VLMOnlyRegistry(vlm))
    state = AgentState(project_id=project["id"], question="What authentication issue is shown?")
    observation = ToolRegistry(state).call(
        "inspect_visual_source",
        source_id="visual-conflict",
        frame_or_image_id=None,
        question_for_vlm=state.question,
    )

    assert observation["conflict_detected"] is True
    assert "visual_conflict_detected" in state.warnings
    assert any(step.get("conflict_detected") is True for step in state.trace)


def test_epic38_structured_llm_orchestrator_accepts_only_bounded_tools(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    llm = FakeLLM(
        [
            {
                "actions": [
                    {"tool": "search_vector_store", "reason": "Need text evidence."},
                    {"tool": "inspect_visual_source", "reason": "Question asks about a diagram."},
                ],
                "rationale": "Use retrieval and visual inspection.",
            }
        ]
    )
    monkeypatch.setattr(agentic_module, "provider_registry", lambda: LLMOnlyRegistry(llm))

    plan = structured_orchestrator_plan("Explain the authentication diagram.", TOOL_NAMES)

    assert [action["tool"] for action in plan["actions"]] == ["search_vector_store", "inspect_visual_source"]
    assert "Allowed tools" in llm.prompts[0]


def test_epic38_unsupported_llm_tool_call_is_rejected(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    llm = FakeLLM([{"actions": [{"tool": "delete_project", "reason": "bad"}], "rationale": "bad"}])
    monkeypatch.setattr(agentic_module, "provider_registry", lambda: LLMOnlyRegistry(llm))

    try:
        structured_orchestrator_plan("Can you answer?", TOOL_NAMES)
    except ValueError as exc:
        assert "Unsupported orchestrator tool" in str(exc)
    else:
        raise AssertionError("Unsupported tool should be rejected")


def test_epic38_structured_llm_critic_can_recommend_ask_vlm(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    llm = FakeLLM(
        [
            {
                "supported_claims": [],
                "unsupported_claims": ["visual claim lacks visual observation"],
                "missing_evidence": ["need frame evidence"],
                "conflicting_evidence": [],
                "citation_quality": "weak",
                "answer_sufficient": False,
                "recommended_action": "ask_vlm_again",
            }
        ]
    )
    monkeypatch.setattr(agentic_module, "provider_registry", lambda: LLMOnlyRegistry(llm))
    evidence = [
        RetrievedChunk(
            id="chunk",
            source_id="source",
            text="Architecture notes mention authentication.",
            citation="source#chunk-1",
            score=0.5,
            metadata={},
        )
    ]

    report = structured_critic_report("What is visible?", "The UI is broken.", [], evidence)

    assert report["recommended_next_action"] == "ask_vlm_again"
    assert report["unsupported_claims_count"] == 1
    assert report["insufficient_evidence"] is True


def test_epic39_old_run_citation_reopens_after_source_delete():
    project = client.post("/api/projects", json={"name": "Snapshot Delete"}).json()
    upload = client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    ).json()
    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()
    citation = run["results"][0]["citations"][0]

    deleted = client.delete(f"/api/sources/{upload['sources'][0]['id']}")
    resolved = client.get(f"/api/citations/resolve?chunk_id={citation['chunk_id']}&run_id={run['id']}")
    reopened = client.get(f"/api/runs/{run['id']}").json()

    assert deleted.status_code == 200
    assert resolved.status_code == 200
    assert resolved.json()["snapshot_restored"] is True
    assert "Authentication uses API Gateway" in resolved.json()["text"]
    assert reopened["knowledge_base_version_status"]["status"] == "archived"
    assert reopened["knowledge_base_version_status"]["snapshot_available"] is True


def test_epic39_clear_replace_marks_old_run_archived_with_snapshot():
    project = client.post("/api/projects", json={"name": "Snapshot Clear"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()

    cleared = client.post(
        f"/api/projects/{project['id']}/sources",
        data={"upload_action": "clear_replace", "confirm_clear": "true"},
        files={"files": ("new.txt", b"New clean benchmark data.", "text/plain")},
    ).json()
    reopened = client.get(f"/api/runs/{run['id']}").json()

    assert cleared["project"]["active_version_id"] != run["knowledge_base_version_id"]
    assert reopened["knowledge_base_version_status"]["status"] == "archived"
    assert reopened["knowledge_base_version_status"]["snapshot_available"] is True


def test_epic39_postgres_is_not_required_product_scope():
    repositories = client.get("/api/settings/repositories").json()
    adapters = client.get("/api/settings/adapters").json()

    assert repositories["selected"] == "sqlite"
    assert repositories["postgres"]["required"] is False
    assert adapters["metadata_stores"]["postgres"]["removed_from_required_scope"] is True


def test_epic41_audio_video_duration_limits_use_ffprobe(monkeypatch):
    monkeypatch.setenv("RAGBENCH_PROVIDER_MODE", "real")
    media_path = data_dir() / "duration.mp4"
    media_path.write_bytes(b"fake media")
    monkeypatch.setattr(main_module, "media_duration_seconds", lambda path: 16 * 60)

    error = main_module.validate_file_limits(media_path, "video")

    assert error == "Video exceeds max_video_duration_minutes=15"
    monkeypatch.setattr(main_module, "media_duration_seconds", lambda path: 31 * 60)
    assert main_module.validate_file_limits(media_path, "audio") == "Audio exceeds max_audio_duration_minutes=30"


def test_epic41_docx_headings_and_tables_preserve_structure_metadata():
    from docx import Document

    project = client.post("/api/projects", json={"name": "DOCX Structure"}).json()
    docx_path = data_dir() / "structured.docx"
    document = Document()
    document.add_heading("Authentication Flow", level=1)
    document.add_paragraph("Login Service issues tokens.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Token Store"
    table.cell(1, 1).text = "Security"
    document.save(docx_path)

    with docx_path.open("rb") as handle:
        upload = client.post(
            f"/api/projects/{project['id']}/sources",
            files={"files": ("structured.docx", handle.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        ).json()
    source_id = upload["sources"][0]["id"]
    content = client.get(f"/api/sources/{source_id}/content").json()
    block_types = [block["block_type"] for block in content["blocks"]]

    assert "heading" in block_types
    assert "table" in block_types
    heading = next(block for block in content["blocks"] if block["block_type"] == "heading")
    assert heading["metadata"]["style"].startswith("Heading")


def test_epic41_docx_structure_budget_limit(monkeypatch):
    from docx import Document

    original = LIMITS["max_docx_pages_estimate"]
    LIMITS["max_docx_pages_estimate"] = 1
    try:
        project = client.post("/api/projects", json={"name": "DOCX Limit"}).json()
        docx_path = data_dir() / "too-large.docx"
        document = Document()
        document.add_paragraph("One")
        document.add_paragraph("Two")
        document.save(docx_path)
        with docx_path.open("rb") as handle:
            upload = client.post(
                f"/api/projects/{project['id']}/sources",
                files={"files": ("too-large.docx", handle.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            ).json()
        assert upload["sources"][0]["status"] == "failed"
        assert "max_docx_pages_estimate" in upload["sources"][0]["error"]
    finally:
        LIMITS["max_docx_pages_estimate"] = original


def test_epic43_model_settings_persist_into_run_config():
    client.patch(
        "/api/settings/models",
        json={
            "llm_model": "qwen3:32b",
            "vlm_model": "qwen3-vl:8b",
            "embedding_model": "bge-m3",
            "transcription_provider": "whisper.cpp",
            "transcription_model": "large-v3-turbo",
            "vector_store": "qdrant",
            "provider_mode": "real",
        },
    )
    project = client.post("/api/projects", json={"name": "Settings Run Config"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )

    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()

    assert run["run_config"]["llm_model"] == "qwen3:32b"
    assert run["run_config"]["vlm_model"] == "qwen3-vl:8b"
    assert run["run_config"]["embedding_model"] == "bge-m3"
    assert run["run_config"]["transcription_provider"] == "whisper.cpp"
    assert run["run_config"]["vector_db_provider"] == "qdrant"
    assert run["run_config"]["prompt_template_versions"]["provider_mode"] == "real"


def test_epic44_sample_dataset_has_all_original_demo_questions_and_audit_traceability():
    sample = client.get("/api/sample-dataset").json()
    audit = client.post("/api/acceptance/no-gap-audit").json()

    assert len(sample["questions"]) == 6
    assert audit["checks"]["post_audit_epics_tracked"] is True
    assert audit["checks"]["real_vlm_captioning_present"] is True
    assert audit["checks"]["source_viewer_content_present"] is True
    assert set(audit["original_plan_sections"]) == {str(index) for index in range(1, 36)}


def test_epic45_feedback_and_result_metrics_are_persisted():
    project = client.post("/api/projects", json={"name": "Feedback Metrics"}).json()
    client.post(
        f"/api/projects/{project['id']}/sources",
        files={"files": ("auth.txt", b"Authentication uses API Gateway and Token Store.", "text/plain")},
    )
    run = client.post(f"/api/projects/{project['id']}/runs", json={"question": "What does authentication use?"}).json()
    saved = client.post(f"/api/runs/{run['id']}/feedback", json={"user_feedback": "Useful comparison."}).json()

    assert saved["user_feedback"] == "Useful comparison."
    for result in saved["results"]:
        assert "prompt_token_estimate" in result["metrics"]
        assert "answer_length" in result["metrics"]
        assert "source_coverage_count" in result["metrics"]


def test_epics46_to_50_final_contract_files_and_audit_checks():
    audit = client.post("/api/acceptance/no-gap-audit").json()
    root = Path(os.environ.get("WORKSPACE_ROOT", "/workspace"))
    project_root = root / "ragbench-studio" if (root / "ragbench-studio").exists() else root
    compose = (project_root / "docker-compose.yml").read_text(encoding="utf-8")
    requirements = (project_root / "apps" / "api" / "requirements.txt").read_text(encoding="utf-8")

    assert audit["checks"]["frontend_stack_decision_documented"] is True
    assert audit["checks"]["prompt_template_files_present"] is True
    assert audit["checks"]["stable_api_export_contracts_present"] is True
    assert audit["checks"]["privacy_non_goals_enforced"] is True
    assert audit["checks"]["single_readme_documentation"] is True
    assert "postgres:" not in compose
    assert "psycopg" not in requirements
