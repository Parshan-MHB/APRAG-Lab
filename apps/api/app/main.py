from __future__ import annotations

import json
import hashlib
import os
import shutil
import time
import uuid
from contextlib import asynccontextmanager
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile, WebSocket
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse, PlainTextResponse, Response, StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader

from .contracts import PROMPT_REGISTRY, default_run_config, prompt_versions_for_pipeline
from .citations import artifact_bytes, resolve_citation, snapshot_run_citations
from .database import connect, data_dir, init_db
from .diagnostics import (
    create_job,
    dependency_status,
    list_job_events,
    list_job_events_by_job,
    local_log_path,
    read_local_logs,
    read_trace_jsonl,
    record_job_event,
    scrub,
    write_local_log,
    write_trace_jsonl,
)
from .guardrails import guardrail_status
from .job_execution import (
    BENCHMARK_STEPS,
    CeleryJobRunner,
    INGESTION_STEPS,
    JOB_STATES,
    LocalJobRunner,
    RedisRQJobRunner,
    cancel_job,
    cancel_run as cancel_run_job,
    create_ingestion_job,
    create_queued_job,
    failure_pipeline_result,
    get_job as load_job,
    get_job_events,
    job_cancel_requested,
    sse_lines,
)
from .migrations import migration_status
from .app_modes import add_message_pair, app_modes_status, conversation_summary, create_session, get_session
from .acceptance import no_gap_audit
from .provider_extensions import extension_settings
from .provider_health import provider_health
from .model_bootstrap import installed_models as ollama_installed_models, pull_model as ollama_pull_model, wait_for_ollama
from .observability import instrument_fastapi, observe_span, observability_status, setup_observability
from .providers import ollama_model_installed, provider_registry, provider_status
from .rag import rebuild_project_graph, run_pipeline, split_chunks, tokenize, recommend, compute_comparison
from .reliability import controlled_timeout_warning, reliability_status
from .resource_profile import preflight_report
from .runtime_adapters import adapter_status
from .runtime_adapters import selected_vector_store
from .media_processing import deterministic_media_enabled, image_blocks, media_duration_seconds, transcribe_audio, video_blocks
from .repositories import repository_status
from .sample_data import DEMO_QUESTIONS, SAMPLE_SOURCES, create_sample_dataset, run_acceptance_suite
from .storage import ensure_storage_layout, parse_trace_jsonl, write_project_manifest, write_run_artifacts
from .workflow import GRAPH_NODES, list_graph_states, persist_graph_state


LIMITS = {
    "max_file_size_mb": 200,
    "max_video_duration_minutes": 15,
    "max_audio_duration_minutes": 30,
    "max_pdf_pages": 300,
    "max_docx_pages_estimate": 300,
    "max_image_count_per_run": 50,
    "max_video_frames_sampled": 30,
    "max_chunks_per_project": 10000,
}

MODEL_SETTINGS_OVERRIDES: dict[str, Any] = {}


@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db()
    yield


setup_observability(os.environ.get("OTEL_SERVICE_NAME", "aprag-lab-api"))
app = FastAPI(title="APRAG-Lab API", lifespan=lifespan)
instrument_fastapi(app)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


class ProjectCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    description: str = ""


class ProjectUpdate(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=120)
    description: str | None = None


class RunCreate(BaseModel):
    question: str = Field(min_length=1, max_length=1000)
    selected_sources: list[str] = []
    run_mode: str = "all_pipelines"
    parent_run_id: str | None = None
    user_feedback: str | None = None


class IngestionJobCreate(BaseModel):
    source_ids: list[str] = []
    fail_steps: list[str] = []


class ModelSettingsPatch(BaseModel):
    llm_model: str | None = None
    vlm_model: str | None = None
    embedding_model: str | None = None
    transcription_provider: str | None = None
    transcription_model: str | None = None
    vector_store: str | None = None
    provider_mode: str | None = None
    provider: str | None = None


class ModelCheckRequest(BaseModel):
    provider: str = "ollama"
    model: str
    model_type: str = "llm"


class ModelPullRequest(BaseModel):
    models: list[str] = []
    project_id: str | None = None


class ConversationSessionCreate(BaseModel):
    title: str = Field(default="Conversation", min_length=1, max_length=120)


class ConversationMessageCreate(BaseModel):
    question: str = Field(min_length=1, max_length=1000)
    parent_run_id: str | None = None
    selected_sources: list[str] = []


class RunFeedbackPatch(BaseModel):
    user_feedback: str = Field(max_length=2000)


class FrontendLogEvent(BaseModel):
    level: str = "info"
    event_type: str
    payload: dict[str, Any] = {}


def now_iso() -> str:
    return datetime.now(UTC).isoformat()


def queue_mode() -> str:
    return os.environ.get("APRAG_QUEUE_MODE", "inline").strip().lower()


def row_to_dict(row: Any) -> dict[str, Any]:
    return dict(row)


def ensure_project_exists(project_id: str) -> None:
    with connect() as conn:
        project = conn.execute("SELECT id FROM projects WHERE id = ?", (project_id,)).fetchone()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")


@app.middleware("http")
async def request_logging_middleware(request: Request, call_next):
    started = time.perf_counter()
    request_id = request.headers.get("x-request-id") or str(uuid.uuid4())
    payload = {
        "request_id": request_id,
        "method": request.method,
        "path": request.url.path,
        "query": str(request.url.query)[:500],
        "client": request.client.host if request.client else "",
        "user_agent": request.headers.get("user-agent", "")[:200],
    }
    write_local_log("api_request_started", payload)
    try:
        response = await call_next(request)
    except Exception as exc:
        elapsed = round(time.perf_counter() - started, 4)
        write_local_log(
            "api_request_failed",
            {
                **payload,
                "elapsed_seconds": elapsed,
                "error_type": type(exc).__name__,
                "error": str(exc),
            },
        )
        raise
    elapsed = round(time.perf_counter() - started, 4)
    response.headers["X-Request-ID"] = request_id
    write_local_log(
        "api_request_completed",
        {
            **payload,
            "status_code": response.status_code,
            "elapsed_seconds": elapsed,
        },
    )
    return response


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "APRAG-Lab-api"}


@app.get("/api/diagnostics/logs")
def diagnostics_logs(limit: int = 200) -> dict[str, Any]:
    return {"path": str(local_log_path()), "entries": read_local_logs(limit)}


@app.post("/api/diagnostics/frontend-event")
def frontend_log_event(payload: FrontendLogEvent) -> dict[str, str]:
    write_local_log(
        "frontend_event",
        {
            "level": payload.level,
            "event_type": payload.event_type,
            **payload.payload,
        },
    )
    return {"status": "logged"}


@app.get("/api/settings/models")
def model_settings() -> dict[str, Any]:
    status = dependency_status()
    resource_profile = preflight_report()
    selected = resource_profile["selected_requirements"]
    models = {key: status[key] for key in ["llm", "vlm", "embeddings", "transcription"]}
    models["overrides"] = MODEL_SETTINGS_OVERRIDES
    models["recommended_models"] = {
        "profile": resource_profile["selected_profile"],
        "llm_model": selected["llm_model"],
        "vlm_model": selected["vlm_model"],
        "embedding_model": selected["embedding_model"],
        "models_to_pull": resource_profile["models_to_pull"],
    }
    models["local_model_warning"] = {
        "provider": "product",
        "available": True,
        "warning": "Local model outputs can be incorrect; answers must be judged by citations and grounding warnings.",
    }
    models["ollama_model_management"] = ollama_model_management_status()
    return models


@app.get("/api/settings/reliability")
def reliability_settings() -> dict[str, Any]:
    return reliability_status()


@app.get("/api/settings/resource-profile")
def resource_profile_settings() -> dict[str, Any]:
    return preflight_report()


@app.get("/api/settings/adapters")
def adapters_settings() -> dict[str, Any]:
    return adapter_status()


@app.get("/api/settings/database-dashboards")
def database_dashboards() -> dict[str, Any]:
    return {
        "dashboards": [
            {
                "id": "qdrant",
                "label": "Qdrant Dashboard",
                "kind": "vector database",
                "url": "http://localhost:6333/dashboard",
                "queryable": True,
                "container": "qdrant",
            },
            {
                "id": "chroma",
                "label": "Chroma API Docs",
                "kind": "vector database",
                "url": "http://localhost:8001/docs",
                "queryable": True,
                "container": "chroma",
            },
            {
                "id": "sqlite",
                "label": "SQLite Query Console",
                "kind": "metadata and graph store",
                "url": "http://localhost:8000/api/database/sqlite-dashboard",
                "queryable": True,
                "container": "api",
            },
            {
                "id": "redis",
                "label": "Redis Commander",
                "kind": "queue/cache store",
                "url": "http://localhost:8083",
                "queryable": True,
                "container": "redis-commander",
            },
        ],
        "opens_in_new_tab": True,
        "warning": "Dashboards expose local development data and should stay bound to localhost.",
    }


@app.get("/api/settings/observability")
def observability_settings() -> dict[str, Any]:
    return observability_status()


def validate_readonly_sql(sql: str) -> str:
    statement = sql.strip().rstrip(";")
    lowered = statement.lower()
    allowed = lowered.startswith("select ") or lowered.startswith("pragma ")
    forbidden = ("insert ", "update ", "delete ", "drop ", "alter ", "create ", "replace ", "attach ", "detach ", "vacuum")
    if not statement or not allowed or any(token in lowered for token in forbidden):
        raise HTTPException(status_code=400, detail="Only read-only SELECT and PRAGMA queries are allowed.")
    return statement


@app.get("/api/database/sqlite-query")
def sqlite_query(sql: str = "SELECT name, type FROM sqlite_master WHERE type IN ('table', 'view') ORDER BY name") -> dict[str, Any]:
    statement = validate_readonly_sql(sql)
    with connect() as conn:
        rows = conn.execute(statement).fetchmany(500)
    return {
        "sql": statement,
        "row_count": len(rows),
        "rows": [dict(row) for row in rows],
        "limit": 500,
    }


@app.get("/api/database/sqlite-dashboard", response_class=HTMLResponse)
def sqlite_dashboard() -> str:
    default_query = "SELECT name, type FROM sqlite_master WHERE type IN ('table', 'view') ORDER BY name"
    return f"""<!doctype html>
<html>
  <head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>SQLite Query Console</title>
    <style>
      body {{ font-family: system-ui, sans-serif; margin: 24px; color: #17202a; background: #f3f5f7; }}
      main {{ max-width: 1100px; margin: 0 auto; }}
      textarea {{ width: 100%; min-height: 120px; font: 14px ui-monospace, SFMono-Regular, Menlo, monospace; }}
      button {{ margin-top: 8px; padding: 8px 12px; }}
      pre {{ background: #fff; border: 1px solid #d7e0e6; border-radius: 6px; overflow: auto; padding: 12px; }}
    </style>
  </head>
  <body>
    <main>
      <h1>SQLite Query Console</h1>
      <p>Read-only local metadata and graph store. Only SELECT and PRAGMA queries are allowed.</p>
      <form id="query-form">
        <textarea id="sql">{default_query}</textarea>
        <br />
        <button type="submit">Run query</button>
      </form>
      <pre id="result">Run a query to inspect the local SQLite database.</pre>
    </main>
    <script>
      const form = document.getElementById('query-form');
      const sql = document.getElementById('sql');
      const result = document.getElementById('result');
      form.addEventListener('submit', async (event) => {{
        event.preventDefault();
        result.textContent = 'Running...';
        const response = await fetch('/api/database/sqlite-query?sql=' + encodeURIComponent(sql.value));
        const payload = await response.json();
        result.textContent = JSON.stringify(payload, null, 2);
      }});
    </script>
  </body>
</html>"""


@app.get("/api/settings/provider-extensions")
def provider_extension_settings() -> dict[str, Any]:
    return extension_settings()


@app.get("/api/app-modes")
def app_modes() -> dict[str, Any]:
    return app_modes_status()


@app.get("/api/settings/migrations")
def migrations_settings() -> dict[str, Any]:
    with connect() as conn:
        return migration_status(conn)


@app.get("/api/settings/repositories")
def repositories_settings() -> dict[str, Any]:
    return repository_status()


@app.get("/api/storage/layout")
def storage_layout() -> dict[str, Any]:
    return {"layout": ensure_storage_layout(), "relative_paths_required": True}


@app.patch("/api/settings/models")
def update_model_settings(payload: ModelSettingsPatch) -> dict[str, Any]:
    updates = payload.model_dump(exclude_none=True)
    MODEL_SETTINGS_OVERRIDES.update(updates)
    return model_settings()


def configured_ollama_model_names() -> list[str]:
    resource_profile = preflight_report()
    selected = resource_profile["selected_requirements"]
    candidates = [
        MODEL_SETTINGS_OVERRIDES.get("llm_model") or selected["llm_model"],
        MODEL_SETTINGS_OVERRIDES.get("vlm_model") or selected["vlm_model"],
        MODEL_SETTINGS_OVERRIDES.get("embedding_model") or selected["embedding_model"],
    ]
    deduped: list[str] = []
    for model in candidates:
        if model and model not in deduped:
            deduped.append(str(model))
    return deduped


def ollama_model_management_status() -> dict[str, Any]:
    resource_profile = preflight_report()
    base_url = os.environ.get("OLLAMA_BASE_URL", "http://host.docker.internal:11434")
    configured = configured_ollama_model_names()
    try:
        installed = sorted(ollama_installed_models(base_url, timeout=1.0))
        reachable = True
        warning = ""
    except Exception as exc:
        installed = []
        reachable = False
        warning = f"Host Ollama is unavailable: {exc}"
    missing = [model for model in configured if not ollama_model_installed(model, installed)]
    return {
        "provider": "ollama",
        "runtime": "host",
        "base_url": base_url,
        "reachable": reachable,
        "configured_models": configured,
        "installed_models": installed,
        "missing_models": missing,
        "safe_to_pull": bool(resource_profile.get("safe_to_pull")),
        "pull_commands": [f"ollama pull {model}" for model in missing],
        "warning": warning,
    }


def model_pull_project_id(preferred_project_id: str | None = None) -> str:
    if preferred_project_id:
        ensure_project_exists(preferred_project_id)
        return preferred_project_id
    with connect() as conn:
        row = conn.execute("SELECT id FROM projects ORDER BY created_at DESC LIMIT 1").fetchone()
    if row:
        return row["id"]
    return create_project(ProjectCreate(name="Local Benchmark Lab"))["id"]


def execute_model_pull_job(job: dict[str, Any]) -> dict[str, Any]:
    base_url = os.environ.get("OLLAMA_BASE_URL", "http://host.docker.internal:11434")
    requested = job["payload"].get("models") or []
    models = [str(model).strip() for model in requested if str(model).strip()]
    if not models:
        models = ollama_model_management_status()["missing_models"]
    installed = set(ollama_installed_models(base_url))
    pulled: list[str] = []
    skipped: list[str] = []
    wait_for_ollama(base_url, attempts=10)
    for model in models:
        if job_cancel_requested(job["id"]):
            record_job_event(job["id"], job["project_id"], "model_pull_cancelled", "cancelled", "Model pull cancelled.", {"pulled": pulled, "skipped": skipped})
            return {"pulled": pulled, "skipped": skipped, "cancelled": True}
        if ollama_model_installed(model, installed):
            skipped.append(model)
            record_job_event(job["id"], job["project_id"], "model_pull_skipped", "running", f"{model} is already installed.", {"model": model})
            continue
        record_job_event(job["id"], job["project_id"], "model_pull_started", "running", f"Pulling {model} into host Ollama.", {"model": model})
        ollama_pull_model(base_url, model)
        installed = set(ollama_installed_models(base_url))
        pulled.append(model)
        record_job_event(job["id"], job["project_id"], "model_pull_finished", "running", f"Pulled {model}.", {"model": model})
    payload = {"pulled": pulled, "skipped": skipped, "status": ollama_model_management_status()}
    record_job_event(job["id"], job["project_id"], "model_pull_complete", "succeeded", "Ollama model pull completed.", payload)
    return payload


@app.get("/api/settings/models/ollama")
def ollama_models_status() -> dict[str, Any]:
    return ollama_model_management_status()


@app.post("/api/settings/models/pull")
def pull_ollama_models(payload: ModelPullRequest) -> dict[str, Any]:
    project_id = model_pull_project_id(payload.project_id)
    models = payload.models or ollama_model_management_status()["missing_models"]
    if not models:
        return {"job_status": "skipped", "models": [], "status": ollama_model_management_status()}
    if queue_mode() != "inline":
        return create_queued_job(project_id, "model_pull", {"models": models})
    job_id = create_job(project_id, status="queued", job_type="model_pull", payload={"models": models})
    job = load_job(job_id)
    record_job_event(job_id, project_id, "queue_worker_started", "running", "Model pull started.", {"models": models})
    result = execute_model_pull_job(job)
    return {"job_id": job_id, "job_status": "succeeded", "models": models, "result": result}


def apply_model_settings_to_run_config(run_config: Any) -> Any:
    if "llm_model" in MODEL_SETTINGS_OVERRIDES:
        run_config.llm_model = MODEL_SETTINGS_OVERRIDES["llm_model"]
    if "vlm_model" in MODEL_SETTINGS_OVERRIDES:
        run_config.vlm_model = MODEL_SETTINGS_OVERRIDES["vlm_model"]
    if "embedding_model" in MODEL_SETTINGS_OVERRIDES:
        run_config.embedding_model = MODEL_SETTINGS_OVERRIDES["embedding_model"]
    if "transcription_provider" in MODEL_SETTINGS_OVERRIDES:
        run_config.transcription_provider = MODEL_SETTINGS_OVERRIDES["transcription_provider"]
    if "transcription_model" in MODEL_SETTINGS_OVERRIDES:
        run_config.transcription_model = MODEL_SETTINGS_OVERRIDES["transcription_model"]
    if "vector_store" in MODEL_SETTINGS_OVERRIDES:
        run_config.vector_db_provider = MODEL_SETTINGS_OVERRIDES["vector_store"]
    if "provider_mode" in MODEL_SETTINGS_OVERRIDES:
        run_config.prompt_template_versions["provider_mode"] = MODEL_SETTINGS_OVERRIDES["provider_mode"]
    return run_config


@app.post("/api/settings/model-check")
def check_model(payload: ModelCheckRequest) -> dict[str, Any]:
    health = provider_health()
    models = health["ollama"].get("models", {})
    installed = models.get(payload.model, {}).get("installed", False)
    return {
        "provider": payload.provider,
        "model": payload.model,
        "model_type": payload.model_type,
        "available": installed,
        "configured": payload.model in str(provider_status()["defaults"]),
        "container_required": False,
        "host_ollama_required": True,
        "progress": controlled_timeout_warning(payload.provider, 10.0),
        "setup": models.get(payload.model, {}).get("remediation") or f"ollama pull {payload.model}",
        "warning": "" if installed else f"Model {payload.model} is missing from host Ollama.",
    }


@app.get("/api/settings/dependencies")
def dependency_settings() -> dict[str, Any]:
    return dependency_status()


@app.get("/api/privacy")
def privacy_status() -> dict[str, Any]:
    return {
        "local_first": True,
        "telemetry_enabled": False,
        "cloud_providers_enabled": False,
        "paid_api_required": False,
        "stored_data": [
            "original uploads",
            "derived chunks",
            "embeddings",
            "transcripts",
            "captions",
            "frames",
            "graph data",
            "trace data",
            "benchmark artifacts",
        ],
        "warnings": ["Local model outputs can be incorrect; verify citations and unsupported-claim warnings."],
    }


@app.get("/api/guardrails")
def roadmap_guardrails() -> dict[str, Any]:
    return guardrail_status()


@app.get("/api/sample-dataset")
def sample_dataset_info() -> dict[str, Any]:
    return {"sources": [source["filename"] for source in SAMPLE_SOURCES], "questions": DEMO_QUESTIONS}


@app.post("/api/sample-dataset/load")
def load_sample_dataset() -> dict[str, Any]:
    sample = create_sample_dataset()
    return {"project": project_summary(sample["project_id"]), "questions": sample["questions"]}


@app.post("/api/acceptance/run")
def acceptance_run() -> dict[str, Any]:
    return run_acceptance_suite()


@app.post("/api/acceptance/no-gap-audit")
def acceptance_no_gap_audit() -> dict[str, Any]:
    return no_gap_audit()


@app.get("/api/settings/providers")
def providers_settings() -> dict[str, Any]:
    return provider_status()


@app.get("/api/settings/provider-health")
def provider_health_settings() -> dict[str, Any]:
    return provider_health()


@app.get("/api/settings/prompts")
def prompt_settings() -> dict[str, Any]:
    return {"prompts": PROMPT_REGISTRY}


@app.post("/api/projects")
def create_project(payload: ProjectCreate) -> dict[str, Any]:
    project_id = str(uuid.uuid4())
    timestamp = now_iso()
    ensure_storage_layout(project_id)
    with connect() as conn:
        conn.execute(
            """
            INSERT INTO projects (id, name, description, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (project_id, payload.name, payload.description, timestamp, timestamp),
        )
    project = get_project(project_id)
    write_project_manifest(project)
    return project


@app.get("/api/projects")
def list_projects() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute("SELECT * FROM projects ORDER BY created_at DESC").fetchall()
    return [project_summary(row["id"]) for row in rows]


@app.get("/api/projects/{project_id}")
def get_project(project_id: str) -> dict[str, Any]:
    return project_summary(project_id)


@app.post("/api/projects/{project_id}/conversation-sessions")
def start_conversation_session(project_id: str, payload: ConversationSessionCreate) -> dict[str, Any]:
    project_summary(project_id)
    return create_session(project_id, payload.title)


@app.get("/api/conversation-sessions/{session_id}")
def read_conversation_session(session_id: str) -> dict[str, Any]:
    try:
        return get_session(session_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Conversation session not found") from None


@app.post("/api/conversation-sessions/{session_id}/messages")
def add_conversation_message(session_id: str, payload: ConversationMessageCreate) -> dict[str, Any]:
    session = read_conversation_session(session_id)
    previous_citations = []
    if payload.parent_run_id:
        parent = get_run(payload.parent_run_id)
        if parent["project_id"] != session["project_id"]:
            raise HTTPException(status_code=400, detail="parent_run_id does not belong to this conversation project")
        for result in parent["results"]:
            previous_citations.extend(result["citations"])
    summary = conversation_summary(payload.question, previous_citations)
    run = create_run(
        session["project_id"],
        RunCreate(
            question=payload.question,
            selected_sources=payload.selected_sources,
            run_mode="conversation_follow_up" if payload.parent_run_id else "conversation",
            parent_run_id=payload.parent_run_id,
            user_feedback=summary,
        ),
    )
    result = next((item for item in run["results"] if item["pipeline_type"] == run["recommendation"]["recommended_flow"]), run["results"][0])
    return add_message_pair(
        session_id,
        session["project_id"],
        payload.question,
        result["answer"],
        run["id"],
        result["citations"],
        summary,
    )


@app.patch("/api/projects/{project_id}")
def update_project(project_id: str, payload: ProjectUpdate) -> dict[str, Any]:
    project_summary(project_id)
    updates: list[str] = []
    values: list[Any] = []
    if payload.name is not None:
        updates.append("name = ?")
        values.append(payload.name)
    if payload.description is not None:
        updates.append("description = ?")
        values.append(payload.description)
    if updates:
        updates.append("updated_at = ?")
        values.append(now_iso())
        values.append(project_id)
        with connect() as conn:
            conn.execute(f"UPDATE projects SET {', '.join(updates)} WHERE id = ?", values)
    project = project_summary(project_id)
    write_project_manifest(project)
    return project


@app.delete("/api/projects/{project_id}")
def delete_project(project_id: str) -> dict[str, str]:
    project_summary(project_id)
    with connect() as conn:
        conn.execute("DELETE FROM projects WHERE id = ?", (project_id,))
    shutil.rmtree(project_dir(project_id), ignore_errors=True)
    return {"status": "deleted", "id": project_id}


def project_summary(project_id: str) -> dict[str, Any]:
    with connect() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        source_count = conn.execute(
            "SELECT COUNT(*) AS count FROM sources WHERE project_id = ?", (project_id,)
        ).fetchone()["count"]
        chunk_count = conn.execute(
            "SELECT COUNT(*) AS count FROM chunks WHERE project_id = ?", (project_id,)
        ).fetchone()["count"]
        run_count = conn.execute(
            "SELECT COUNT(*) AS count FROM runs WHERE project_id = ?", (project_id,)
        ).fetchone()["count"]
        active_version = None
        if project["active_version_id"]:
            version = conn.execute(
                "SELECT * FROM knowledge_base_versions WHERE id = ?",
                (project["active_version_id"],),
            ).fetchone()
            active_version = version_to_dict(version) if version else None
    data = row_to_dict(project)
    data.update(
        {
            "source_count": source_count,
            "chunk_count": chunk_count,
            "run_count": run_count,
            "active_version": active_version,
        }
    )
    return data


def version_to_dict(row: Any) -> dict[str, Any]:
    data = row_to_dict(row)
    data["source_ids"] = json.loads(data.pop("source_ids_json"))
    return data


def latest_version_number(conn: Any, project_id: str) -> int:
    row = conn.execute(
        "SELECT COALESCE(MAX(version_number), 0) AS version_number FROM knowledge_base_versions WHERE project_id = ?",
        (project_id,),
    ).fetchone()
    return int(row["version_number"])


def create_kb_version(
    conn: Any,
    project_id: str,
    change_type: str,
    source_ids: list[str],
    notes: str = "",
) -> str:
    version_id = str(uuid.uuid4())
    version_number = latest_version_number(conn, project_id) + 1
    chunk_count = conn.execute(
        "SELECT COUNT(*) AS count FROM chunks WHERE project_id = ?",
        (project_id,),
    ).fetchone()["count"]
    conn.execute(
        """
        INSERT INTO knowledge_base_versions
          (id, project_id, version_number, created_at, change_type, source_ids_json, chunk_count, embedding_model, graph_version, notes)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            version_id,
            project_id,
            version_number,
            now_iso(),
            change_type,
            json.dumps(source_ids),
            chunk_count,
            provider_registry().embeddings().model,
            str(version_number),
            notes,
        ),
    )
    conn.execute(
        """
        UPDATE projects
        SET active_version_id = ?, processing_status = ?, graph_status = ?, updated_at = ?
        WHERE id = ?
        """,
        (version_id, "processed", "ready", now_iso(), project_id),
    )
    return version_id


@app.get("/api/projects/{project_id}/versions")
def list_versions(project_id: str) -> list[dict[str, Any]]:
    project_summary(project_id)
    with connect() as conn:
        rows = conn.execute(
            "SELECT * FROM knowledge_base_versions WHERE project_id = ? ORDER BY version_number DESC",
            (project_id,),
        ).fetchall()
    return [version_to_dict(row) for row in rows]


def project_dir(project_id: str) -> Path:
    return data_dir() / "projects" / project_id


def source_type(filename: str, content_type: str | None) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".csv", ".json"}:
        return "text"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        return "image"
    if suffix in {".wav", ".mp3", ".m4a", ".ogg"}:
        return "audio"
    if suffix in {".mp4", ".mov", ".mkv", ".webm"}:
        return "video"
    if content_type and content_type.startswith("text/"):
        return "text"
    return "unsupported"


def deterministic_vector(text: str, size: int = 16) -> list[float]:
    vector = [0.0] * size
    terms = tokenize(text)
    for term in terms:
        digest = hashlib.sha256(term.encode("utf-8")).digest()
        vector[digest[0] % size] += 1.0
    total = sum(vector) or 1.0
    return [round(value / total, 6) for value in vector]


def derived_dir(project_id: str, name: str) -> Path:
    path = project_dir(project_id) / "sources" / "derived" / name
    path.mkdir(parents=True, exist_ok=True)
    return path


def validate_file_limits(path: Path, kind: str) -> str | None:
    max_bytes = LIMITS["max_file_size_mb"] * 1024 * 1024
    if path.stat().st_size > max_bytes:
        return f"File exceeds max_file_size_mb={LIMITS['max_file_size_mb']}"
    if kind in {"audio", "video"} and not deterministic_media_enabled():
        try:
            duration = media_duration_seconds(path)
        except Exception as exc:
            return f"Could not inspect media duration with FFprobe: {exc}"
        if duration is not None:
            limit_minutes = LIMITS["max_audio_duration_minutes"] if kind == "audio" else LIMITS["max_video_duration_minutes"]
            if duration > limit_minutes * 60:
                return f"{kind.title()} exceeds max_{kind}_duration_minutes={limit_minutes}"
    return None


def persist_blocks_for_source(conn: Any, project_id: str, source_id: str, filename: str, kind: str, blocks: list[dict[str, Any]]) -> None:
    for block in blocks:
        content_block_id = str(uuid.uuid4())
        metadata = block.get("metadata", {})
        conn.execute(
            """
            INSERT INTO content_blocks
              (id, project_id, source_id, block_type, text, page_number, timestamp_start, timestamp_end,
               frame_path, image_region, confidence, metadata_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                content_block_id,
                project_id,
                source_id,
                block["block_type"],
                block["text"],
                block.get("page_number"),
                block.get("timestamp_start"),
                block.get("timestamp_end"),
                block.get("frame_path"),
                block.get("image_region"),
                block.get("confidence"),
                json.dumps(metadata),
            ),
        )
        for index, chunk in enumerate(split_chunks(block["text"])):
            citation_suffix = f"chunk-{index + 1}"
            if block.get("page_number"):
                citation_suffix = f"page-{block['page_number']}-{citation_suffix}"
            elif block.get("timestamp_start") is not None:
                citation_suffix = f"t-{block.get('timestamp_start', 0)}-{citation_suffix}"
            source_reference_label = f"{filename}#{citation_suffix}"
            chunk_id = str(uuid.uuid4())
            embedding_id = str(uuid.uuid4())
            embedding_adapter = provider_registry().embeddings()
            embedding_vector = embedding_adapter.embed_text(chunk)
            chunk_metadata = {
                "chunk_index": index + 1,
                "source_type": kind,
                "block_type": block["block_type"],
                "embedding_provider": embedding_adapter.provider,
                "embedding_model": embedding_adapter.model,
                **metadata,
            }
            conn.execute(
                """
                INSERT INTO chunks
                  (id, project_id, source_id, content_block_id, text, citation, token_estimate,
                   embedding_id, page_number, timestamp_start, timestamp_end, source_reference_label, metadata_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    chunk_id,
                    project_id,
                    source_id,
                    content_block_id,
                    chunk,
                    source_reference_label,
                    len(tokenize(chunk)),
                    embedding_id,
                    block.get("page_number"),
                    block.get("timestamp_start"),
                    block.get("timestamp_end"),
                    source_reference_label,
                    json.dumps(chunk_metadata),
                ),
            )
            conn.execute(
                """
                INSERT INTO vector_records
                  (id, project_id, source_id, chunk_id, embedding_model, vector_json, metadata_json)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    embedding_id,
                    project_id,
                    source_id,
                    chunk_id,
                    embedding_adapter.model,
                    json.dumps(embedding_vector),
                    json.dumps({"source_reference_label": source_reference_label}),
                ),
            )
            selected_vector_store(project_id).upsert(
                [
                    {
                        "id": embedding_id,
                        "project_id": project_id,
                        "source_id": source_id,
                        "chunk_id": chunk_id,
                        "text": chunk,
                        "vector": embedding_vector,
                        "metadata": {
                            "source_id": source_id,
                            "chunk_id": chunk_id,
                            "citation": source_reference_label,
                            "source_reference_label": source_reference_label,
                            "source_type": kind,
                            "page_number": block.get("page_number"),
                            "timestamp_start": block.get("timestamp_start"),
                            "timestamp_end": block.get("timestamp_end"),
                        },
                    }
                ]
            )


def process_staged_sources(project_id: str, source_ids: list[str], change_type: str = "append_sources") -> dict[str, Any]:
    uploaded: list[dict[str, Any]] = []
    print(f"APRAG-Lab source ingestion started project_id={project_id} sources={len(source_ids)} change_type={change_type}", flush=True)
    with connect() as conn:
        conn.execute("UPDATE projects SET processing_status = ?, updated_at = ? WHERE id = ?", ("processing", now_iso(), project_id))
        rows = conn.execute(
            "SELECT * FROM sources WHERE project_id = ? AND id IN ({})".format(",".join("?" for _ in source_ids)),
            [project_id, *source_ids],
        ).fetchall() if source_ids else []
    for source in rows:
        target = data_dir() / source["local_path"]
        kind = source["source_type"]
        print(f"APRAG-Lab source processing started source_id={source['id']} filename={source['filename']} type={kind}", flush=True)
        limit_error = validate_file_limits(target, kind)
        blocks, extraction_error = ([], limit_error) if limit_error else extract_content_blocks(target, kind, source["filename"], project_id, source["id"])
        status = "processed"
        error = extraction_error
        if kind == "unsupported":
            status = "failed"
            error = "Unsupported source type"
        elif not blocks or not any(block.get("text", "").strip() for block in blocks):
            status = "failed"
            error = error or "No text could be extracted"
        with connect() as conn:
            conn.execute("UPDATE sources SET status = ? WHERE id = ?", ("processing", source["id"]))
            conn.execute("UPDATE sources SET status = ?, error = ? WHERE id = ?", (status, error, source["id"]))
            if status == "processed":
                selected_vector_store(project_id).delete_source(source["id"])
                conn.execute("DELETE FROM vector_records WHERE source_id = ?", (source["id"],))
                conn.execute("DELETE FROM content_blocks WHERE source_id = ?", (source["id"],))
                conn.execute("DELETE FROM chunks WHERE source_id = ?", (source["id"],))
                persist_blocks_for_source(conn, project_id, source["id"], source["filename"], kind, blocks)
        uploaded.append({"id": source["id"], "filename": source["filename"], "source_type": kind, "status": status, "error": error})
        print(f"APRAG-Lab source processing finished source_id={source['id']} status={status} error={error or ''}", flush=True)
    rebuild_project_graph(project_id)
    with connect() as conn:
        processed_count = conn.execute("SELECT COUNT(*) AS count FROM sources WHERE project_id = ? AND status = 'processed'", (project_id,)).fetchone()["count"]
        failed_count = conn.execute("SELECT COUNT(*) AS count FROM sources WHERE project_id = ? AND status = 'failed'", (project_id,)).fetchone()["count"]
        version_id = create_kb_version(conn, project_id, change_type, source_ids)
        if source_ids:
            conn.execute(
                "UPDATE sources SET knowledge_base_version_id = ? WHERE id IN ({})".format(",".join("?" for _ in source_ids)),
                [version_id, *source_ids],
            )
        if processed_count and failed_count:
            conn.execute("UPDATE projects SET processing_status = ? WHERE id = ?", ("partial_success", project_id))
    project = project_summary(project_id)
    write_project_manifest(project)
    print(f"APRAG-Lab source ingestion finished project_id={project_id} processed={processed_count} failed={failed_count}", flush=True)
    return {"sources": uploaded, "project": project, "target_project_id": project_id, "upload_action": change_type}


def extract_content_blocks(path: Path, kind: str, filename: str, project_id: str, source_id: str) -> tuple[list[dict[str, Any]], str | None]:
    try:
        project_root = project_dir(project_id)
        if kind == "text":
            return [
                {
                    "block_type": "text",
                    "text": path.read_text(encoding="utf-8", errors="replace"),
                    "metadata": {"filename": filename},
                }
            ], None
        if kind == "pdf":
            reader = PdfReader(str(path))
            if len(reader.pages) > LIMITS["max_pdf_pages"]:
                return [], f"PDF exceeds max_pdf_pages={LIMITS['max_pdf_pages']}"
            return [
                {
                    "block_type": "page_text",
                    "text": page.extract_text() or "",
                    "page_number": index + 1,
                    "metadata": {"filename": filename},
                }
                for index, page in enumerate(reader.pages)
            ], None
        if kind == "docx":
            document = Document(str(path))
            structural_units = len(document.paragraphs) + sum(len(table.rows) for table in document.tables)
            if structural_units > LIMITS["max_docx_pages_estimate"]:
                return [], f"DOCX exceeds max_docx_pages_estimate={LIMITS['max_docx_pages_estimate']}"
            blocks: list[dict[str, Any]] = []
            for index, paragraph in enumerate(document.paragraphs):
                if paragraph.text.strip():
                    style_name = paragraph.style.name if paragraph.style else ""
                    blocks.append(
                        {
                            "block_type": "heading" if style_name.lower().startswith("heading") else "text",
                            "text": paragraph.text,
                            "metadata": {"filename": filename, "paragraph_index": index, "style": style_name},
                        }
                    )
            for table_index, table in enumerate(document.tables):
                rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
                if rows:
                    blocks.append(
                        {
                            "block_type": "table",
                            "text": "\n".join(rows),
                            "metadata": {"filename": filename, "table_index": table_index},
                        }
                    )
            return blocks, None
        if kind == "image":
            blocks, warnings = image_blocks(path, filename, project_root, source_id)
            return blocks, "; ".join(warnings) if warnings and not blocks else None
        if kind == "audio":
            blocks, warnings = transcribe_audio(path, project_root, source_id)
            return blocks, "; ".join(warnings) if warnings and not blocks else None
        if kind == "video":
            blocks, warnings = video_blocks(path, filename, project_root, source_id)
            return blocks, "; ".join(warnings) if warnings and not blocks else None
        return [], "unsupported_source_type"
    except Exception as exc:
        return [], str(exc)


@app.post("/api/projects/{project_id}/sources")
def upload_sources(
    project_id: str,
    files: list[UploadFile] = File(...),
    upload_action: str = Form("append"),
    confirm_clear: bool = Form(False),
    new_project_name: str | None = Form(None),
) -> dict[str, Any]:
    project_summary(project_id)
    if upload_action not in {"append", "clear_replace", "create_new"}:
        raise HTTPException(status_code=400, detail="upload_action must be append, clear_replace, or create_new")
    if upload_action != "create_new":
        ensure_no_active_source_upload(project_id)

    with connect() as conn:
        existing_source_count = conn.execute(
            "SELECT COUNT(*) AS count FROM sources WHERE project_id = ?",
            (project_id,),
        ).fetchone()["count"]

    target_project_id = project_id
    if upload_action == "create_new":
        created = create_project(
            ProjectCreate(
                name=new_project_name or f"{project_summary(project_id)['name']} Copy",
                description="Created from upload choice workflow.",
            )
        )
        target_project_id = created["id"]
    elif upload_action == "clear_replace" and existing_source_count > 0:
        if not confirm_clear:
            raise HTTPException(status_code=400, detail="confirm_clear is required for clear and replace")
        clear_project_sources(project_id)

    change_type = "initial_upload"
    if upload_action == "append" and existing_source_count > 0:
        change_type = "append_sources"
    elif upload_action == "clear_replace":
        change_type = "clear_and_replace"

    original_dir = project_dir(project_id) / "sources" / "original"
    if target_project_id != project_id:
        original_dir = project_dir(target_project_id) / "sources" / "original"
    original_dir.mkdir(parents=True, exist_ok=True)
    uploaded: list[dict[str, Any]] = []
    uploaded_source_ids: list[str] = []
    image_upload_count = sum(1 for upload in files if source_type(upload.filename or "upload", upload.content_type) == "image")
    if image_upload_count > LIMITS["max_image_count_per_run"]:
        raise HTTPException(status_code=400, detail=f"Upload exceeds max_image_count_per_run={LIMITS['max_image_count_per_run']}")

    for upload in files:
        kind = source_type(upload.filename or "upload", upload.content_type)
        source_id = str(uuid.uuid4())
        target = original_dir / f"{source_id}_{Path(upload.filename or 'upload').name}"
        with target.open("wb") as handle:
            shutil.copyfileobj(upload.file, handle)

        if queue_mode() != "inline":
            timestamp = now_iso()
            with connect() as conn:
                conn.execute(
                    """
                    INSERT INTO sources
                      (id, project_id, knowledge_base_version_id, filename, source_type, mime_type, local_path, status, error, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        source_id,
                        target_project_id,
                        None,
                        upload.filename or target.name,
                        kind,
                        upload.content_type or "",
                        str(target.relative_to(data_dir())),
                        "queued",
                        None,
                        timestamp,
                    ),
                )
                conn.execute("UPDATE projects SET updated_at = ?, processing_status = ? WHERE id = ?", (timestamp, "processing", target_project_id))
            uploaded.append({"id": source_id, "filename": upload.filename, "source_type": kind, "status": "queued", "error": None})
            uploaded_source_ids.append(source_id)
            continue

        limit_error = validate_file_limits(target, kind)
        blocks, extraction_error = ([], limit_error) if limit_error else extract_content_blocks(
            target,
            kind,
            upload.filename or target.name,
            target_project_id,
            source_id,
        )
        status = "processed"
        error = extraction_error
        if kind == "unsupported":
            status = "failed"
            error = "Unsupported source type"
        elif not blocks or not any(block.get("text", "").strip() for block in blocks):
            status = "failed"
            error = error or "No text could be extracted"

        timestamp = now_iso()
        with connect() as conn:
            conn.execute(
                """
                INSERT INTO sources
                  (id, project_id, knowledge_base_version_id, filename, source_type, mime_type, local_path, status, error, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    source_id,
                    target_project_id,
                    None,
                    upload.filename or target.name,
                    kind,
                    upload.content_type or "",
                    str(target.relative_to(data_dir())),
                    status,
                    error,
                    timestamp,
                ),
            )
            if status == "processed":
                for block in blocks:
                    content_block_id = str(uuid.uuid4())
                    metadata = block.get("metadata", {})
                    conn.execute(
                        """
                        INSERT INTO content_blocks
                          (id, project_id, source_id, block_type, text, page_number, timestamp_start, timestamp_end,
                           frame_path, image_region, confidence, metadata_json)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            content_block_id,
                            target_project_id,
                            source_id,
                            block["block_type"],
                            block["text"],
                            block.get("page_number"),
                            block.get("timestamp_start"),
                            block.get("timestamp_end"),
                            block.get("frame_path"),
                            block.get("image_region"),
                            block.get("confidence"),
                            json.dumps(metadata),
                        ),
                    )
                    for index, chunk in enumerate(split_chunks(block["text"])):
                        citation_suffix = f"chunk-{index + 1}"
                        if block.get("page_number"):
                            citation_suffix = f"page-{block['page_number']}-{citation_suffix}"
                        elif block.get("timestamp_start") is not None:
                            citation_suffix = f"t-{block.get('timestamp_start', 0)}-{citation_suffix}"
                        source_reference_label = f"{upload.filename or target.name}#{citation_suffix}"
                        chunk_id = str(uuid.uuid4())
                        embedding_id = str(uuid.uuid4())
                        embedding_adapter = provider_registry().embeddings()
                        embedding_vector = embedding_adapter.embed_text(chunk)
                        chunk_metadata = {
                            "chunk_index": index + 1,
                            "source_type": kind,
                            "block_type": block["block_type"],
                            "embedding_provider": embedding_adapter.provider,
                            "embedding_model": embedding_adapter.model,
                            **metadata,
                        }
                        conn.execute(
                            """
                            INSERT INTO chunks
                              (id, project_id, source_id, content_block_id, text, citation, token_estimate,
                               embedding_id, page_number, timestamp_start, timestamp_end, source_reference_label, metadata_json)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                            """,
                            (
                                chunk_id,
                                target_project_id,
                                source_id,
                                content_block_id,
                                chunk,
                                source_reference_label,
                                len(tokenize(chunk)),
                                embedding_id,
                                block.get("page_number"),
                                block.get("timestamp_start"),
                                block.get("timestamp_end"),
                                source_reference_label,
                                json.dumps(chunk_metadata),
                            ),
                        )
                        conn.execute(
                            """
                            INSERT INTO vector_records
                              (id, project_id, source_id, chunk_id, embedding_model, vector_json, metadata_json)
                            VALUES (?, ?, ?, ?, ?, ?, ?)
                            """,
                            (
                                embedding_id,
                                target_project_id,
                                source_id,
                                chunk_id,
                                embedding_adapter.model,
                                json.dumps(embedding_vector),
                                json.dumps({"source_reference_label": source_reference_label}),
                            ),
                        )
                        selected_vector_store(target_project_id).upsert(
                            [
                                {
                                    "id": embedding_id,
                                    "project_id": target_project_id,
                                    "source_id": source_id,
                                    "chunk_id": chunk_id,
                                    "text": chunk,
                                    "vector": embedding_vector,
                                    "metadata": {
                                        "source_id": source_id,
                                        "chunk_id": chunk_id,
                                        "citation": source_reference_label,
                                        "source_reference_label": source_reference_label,
                                        "source_type": kind,
                                        "page_number": block.get("page_number"),
                                        "timestamp_start": block.get("timestamp_start"),
                                        "timestamp_end": block.get("timestamp_end"),
                                    },
                                }
                            ]
                        )
                project_chunk_count = conn.execute(
                    "SELECT COUNT(*) AS count FROM chunks WHERE project_id = ?",
                    (target_project_id,),
                ).fetchone()["count"]
                if project_chunk_count > LIMITS["max_chunks_per_project"]:
                    raise HTTPException(status_code=400, detail=f"Project exceeds max_chunks_per_project={LIMITS['max_chunks_per_project']}")
            conn.execute("UPDATE projects SET updated_at = ?, processing_status = ? WHERE id = ?", (timestamp, "processed", target_project_id))
        uploaded.append({"id": source_id, "filename": upload.filename, "source_type": kind, "status": status, "error": error})
        uploaded_source_ids.append(source_id)

    if queue_mode() != "inline":
        job = create_queued_job(
            target_project_id,
            "source_upload",
            {"source_ids": uploaded_source_ids, "change_type": change_type},
        )
        project = project_summary(target_project_id)
        return {
            "sources": uploaded,
            "project": project,
            "upload_action": upload_action,
            "target_project_id": target_project_id,
            "job_id": job["id"],
            "job_status": job["status"],
        }

    rebuild_project_graph(target_project_id)
    with connect() as conn:
        processed_count = conn.execute(
            "SELECT COUNT(*) AS count FROM sources WHERE project_id = ? AND status = 'processed'",
            (target_project_id,),
        ).fetchone()["count"]
        failed_count = conn.execute(
            "SELECT COUNT(*) AS count FROM sources WHERE project_id = ? AND status = 'failed'",
            (target_project_id,),
        ).fetchone()["count"]
        version_id = create_kb_version(conn, target_project_id, change_type, uploaded_source_ids)
        conn.execute(
            "UPDATE sources SET knowledge_base_version_id = ? WHERE id IN ({})".format(
                ",".join("?" for _ in uploaded_source_ids)
            ),
            [version_id, *uploaded_source_ids],
        )
        if processed_count and failed_count:
            conn.execute(
                "UPDATE projects SET processing_status = ? WHERE id = ?",
                ("partial_success", target_project_id),
            )
    project = project_summary(target_project_id)
    write_project_manifest(project)
    return {
        "sources": uploaded,
        "project": project,
        "upload_action": upload_action,
        "target_project_id": target_project_id,
    }


def ensure_no_active_source_upload(project_id: str) -> None:
    with connect() as conn:
        active = conn.execute(
            """
            SELECT id, status
            FROM jobs
            WHERE project_id = ?
              AND job_type = 'source_upload'
              AND status IN ('queued', 'running')
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (project_id,),
        ).fetchone()
    if active:
        raise HTTPException(
            status_code=409,
            detail=f"Source upload job {active['id']} is still {active['status']}; wait for it to finish before changing this knowledge base.",
        )


def ensure_project_ready_for_benchmark(project: dict[str, Any]) -> None:
    if project.get("processing_status") in {"queued", "processing", "running"}:
        raise HTTPException(
            status_code=409,
            detail="Source extraction is still processing. Wait until the extraction queue is complete before running a benchmark.",
        )
    ensure_no_active_source_upload(project["id"])
    if not project.get("active_version_id"):
        raise HTTPException(
            status_code=400,
            detail="Upload and process at least one source before running a benchmark",
        )


def clear_project_sources(project_id: str) -> None:
    selected_vector_store(project_id).delete_project()
    with connect() as conn:
        conn.execute("DELETE FROM relationships WHERE project_id = ?", (project_id,))
        conn.execute("DELETE FROM entities WHERE project_id = ?", (project_id,))
        conn.execute("DELETE FROM vector_records WHERE project_id = ?", (project_id,))
        conn.execute("DELETE FROM chunks WHERE project_id = ?", (project_id,))
        conn.execute("DELETE FROM content_blocks WHERE project_id = ?", (project_id,))
        conn.execute("DELETE FROM sources WHERE project_id = ?", (project_id,))
        conn.execute(
            "UPDATE projects SET processing_status = ?, graph_status = ?, updated_at = ? WHERE id = ?",
            ("empty", "empty", now_iso(), project_id),
        )
    shutil.rmtree(project_dir(project_id) / "sources", ignore_errors=True)


@app.get("/api/projects/{project_id}/sources")
def list_sources(project_id: str) -> list[dict[str, Any]]:
    project_summary(project_id)
    with connect() as conn:
        rows = conn.execute(
            "SELECT * FROM sources WHERE project_id = ? ORDER BY created_at DESC",
            (project_id,),
        ).fetchall()
    return [row_to_dict(row) for row in rows]


@app.post("/api/projects/{project_id}/ingestion-jobs")
def start_ingestion_job(project_id: str, payload: IngestionJobCreate | None = None) -> dict[str, Any]:
    project_summary(project_id)
    payload = payload or IngestionJobCreate()
    invalid_steps = sorted(set(payload.fail_steps) - set(INGESTION_STEPS))
    if invalid_steps:
        raise HTTPException(status_code=400, detail=f"Unknown ingestion steps: {', '.join(invalid_steps)}")
    if queue_mode() != "inline":
        return create_queued_job(
            project_id,
            "ingestion",
            {"source_ids": payload.source_ids, "fail_steps": payload.fail_steps},
        )
    return create_ingestion_job(project_id, payload.source_ids, payload.fail_steps)


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str) -> dict[str, Any]:
    try:
        return load_job(job_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Job not found") from None


@app.get("/api/jobs/{job_id}/events")
def job_events(job_id: str) -> list[dict[str, Any]]:
    try:
        return get_job_events(job_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Job not found") from None


@app.get("/api/jobs/{job_id}/events.sse")
def job_events_sse(job_id: str) -> StreamingResponse:
    return StreamingResponse(iter([sse_lines(job_events(job_id))]), media_type="text/event-stream")


@app.websocket("/api/jobs/{job_id}/events.ws")
async def job_events_ws(websocket: WebSocket, job_id: str) -> None:
    await websocket.accept()
    try:
        for event in job_events(job_id):
            await websocket.send_json(event)
    except HTTPException:
        await websocket.send_json({"error": "Job not found"})
    await websocket.close()


@app.get("/api/sources/{source_id}")
def get_source(source_id: str) -> dict[str, Any]:
    with connect() as conn:
        source = conn.execute("SELECT * FROM sources WHERE id = ?", (source_id,)).fetchone()
    if not source:
        raise HTTPException(status_code=404, detail="Source not found")
    return row_to_dict(source)


@app.get("/api/sources/{source_id}/content")
def get_source_content(source_id: str) -> dict[str, Any]:
    source = get_source(source_id)
    with connect() as conn:
        blocks = conn.execute(
            "SELECT * FROM content_blocks WHERE source_id = ? ORDER BY page_number, timestamp_start, id",
            (source_id,),
        ).fetchall()
    return {
        "source": source,
        "blocks": [
            {
                **row_to_dict(block),
                "metadata": json.loads(block["metadata_json"]),
            }
            for block in blocks
        ],
    }


@app.get("/api/citations/resolve")
def citation_resolve(chunk_id: str | None = None, label: str | None = None, run_id: str | None = None) -> dict[str, Any]:
    resolution = resolve_citation(chunk_id=chunk_id, label=label, run_id=run_id)
    if not resolution:
        raise HTTPException(status_code=404, detail="Citation could not be resolved")
    return resolution


@app.get("/api/artifacts/{artifact_path:path}")
def read_artifact(artifact_path: str) -> Response:
    try:
        content, media_type = artifact_bytes(artifact_path)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Artifact not found") from None
    return Response(content=content, media_type=media_type)


@app.delete("/api/sources/{source_id}")
def delete_source(source_id: str) -> dict[str, Any]:
    source = get_source(source_id)
    project_id = source["project_id"]
    selected_vector_store(project_id).delete_source(source_id)
    local_path = data_dir() / source["local_path"]
    with connect() as conn:
        conn.execute("DELETE FROM vector_records WHERE source_id = ?", (source_id,))
        conn.execute("DELETE FROM content_blocks WHERE source_id = ?", (source_id,))
        conn.execute("DELETE FROM sources WHERE id = ?", (source_id,))
        conn.execute("DELETE FROM chunks WHERE source_id = ?", (source_id,))
        create_kb_version(conn, project_id, "remove_sources", [source_id], notes=f"Removed source {source['filename']}")
    if local_path.exists():
        local_path.unlink()
    rebuild_project_graph(project_id)
    return {"status": "deleted", "source_id": source_id, "project": project_summary(project_id)}


def run_pipeline_safely(project_id: str, question: str, pipeline: str, selected_source_ids: list[str]) -> dict[str, Any]:
    with observe_span(
        "rag.pipeline.run",
        {
            "project_id": project_id,
            "pipeline_type": pipeline,
            "selected_source_count": len(selected_source_ids),
            "question_length": len(question),
        },
        {"question": question, "selected_sources": selected_source_ids},
    ) as span:
        try:
            result = run_pipeline(project_id, question, pipeline, selected_source_ids=selected_source_ids)
            if span:
                span.set_attribute("citation_count", len(result.get("citations", [])))
                span.set_attribute("trace_step_count", len(result.get("trace", [])))
                span.set_attribute("warning_count", len(result.get("warnings", [])))
            return result
        except Exception as exc:
            if span:
                span.set_attribute("pipeline_failed", True)
            return failure_pipeline_result(pipeline, exc)


@app.post("/api/projects/{project_id}/runs")
def create_run(project_id: str, payload: RunCreate) -> dict[str, Any]:
    project = project_summary(project_id)
    ensure_project_ready_for_benchmark(project)
    with connect() as conn:
        chunk_count = conn.execute(
            "SELECT COUNT(*) AS count FROM chunks WHERE project_id = ?",
            (project_id,),
        ).fetchone()["count"]
        if payload.parent_run_id:
            parent = conn.execute(
                "SELECT id FROM runs WHERE id = ? AND project_id = ?",
                (payload.parent_run_id, project_id),
            ).fetchone()
            if not parent:
                raise HTTPException(status_code=400, detail="parent_run_id does not belong to this project")
    if chunk_count == 0:
        raise HTTPException(status_code=400, detail="Upload and process at least one source before running a benchmark")

    run_id = str(uuid.uuid4())
    timestamp = now_iso()
    run_config = apply_model_settings_to_run_config(default_run_config(project["active_version_id"]))
    if queue_mode() != "inline":
        with connect() as conn:
            conn.execute(
                """
                INSERT INTO runs
                  (id, project_id, knowledge_base_version_id, question, selected_sources_json, run_mode,
                   recommendation_json, parent_run_id, user_feedback, run_config_json, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id,
                    project_id,
                    project["active_version_id"],
                    payload.question,
                    json.dumps(payload.selected_sources),
                    payload.run_mode,
                    json.dumps({"recommended_flow": "queued", "reasons": ["Run is queued for worker execution."], "tradeoff_notes": []}),
                    payload.parent_run_id,
                    payload.user_feedback,
                    run_config.model_dump_json(),
                    timestamp,
                ),
            )
        job = create_queued_job(
            project_id,
            "benchmark",
            {"run_id": run_id, "question": payload.question, "selected_sources": payload.selected_sources},
            run_id=run_id,
        )
        run = get_run(run_id)
        run["job_id"] = job["id"]
        run["job_status"] = job["status"]
        return run

    job_id = create_job(project_id, run_id=run_id, status="queued", job_type="benchmark", payload={"question": payload.question})
    record_job_event(job_id, project_id, "create_run", "running", "Benchmark run accepted.", {"question": payload.question, "selected_sources": payload.selected_sources}, run_id)
    record_job_event(job_id, project_id, "start_traditional_pipeline", "running", "Traditional pipeline started.", run_id=run_id)
    traditional = run_pipeline_safely(project_id, payload.question, "traditional", payload.selected_sources)
    record_job_event(job_id, project_id, "start_agentic_pipeline", "running", "Agentic pipeline started.", run_id=run_id)
    agentic = run_pipeline_safely(project_id, payload.question, "agentic", payload.selected_sources)
    record_job_event(job_id, project_id, "start_hybrid_graph_pipeline", "running", "Hybrid Graph pipeline started.", run_id=run_id)
    hybrid = run_pipeline_safely(project_id, payload.question, "hybrid_graph", payload.selected_sources)
    results = [traditional, agentic, hybrid]
    failed_pipelines = [result["pipeline_type"] for result in results if "pipeline_failed" in result.get("warnings", [])]
    record_job_event(
        job_id,
        project_id,
        "collect_results",
        "running",
        "Pipeline results collected.",
        {"failed_pipelines": failed_pipelines},
        run_id,
    )
    recommendation = recommend(results)
    record_job_event(job_id, project_id, "compute_comparison", "running", "Recommendation and comparison computed.", {"recommended_flow": recommendation["recommended_flow"]}, run_id)

    with connect() as conn:
        conn.execute(
            """
            INSERT INTO runs
              (id, project_id, knowledge_base_version_id, question, selected_sources_json, run_mode,
               recommendation_json, parent_run_id, user_feedback, run_config_json, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                project_id,
                project["active_version_id"],
                payload.question,
                json.dumps(payload.selected_sources),
                payload.run_mode,
                json.dumps(recommendation),
                payload.parent_run_id,
                payload.user_feedback,
                run_config.model_dump_json(),
                timestamp,
            ),
        )
        result_ids: dict[str, str] = {}
        for result in results:
            result_id = str(uuid.uuid4())
            result_ids[result["pipeline_type"]] = result_id
            conn.execute(
                """
                INSERT INTO pipeline_results
                  (id, run_id, pipeline_type, answer, citations_json, trace_json, metrics_json, techniques_json, warnings_json, prompt_versions_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    result_id,
                    run_id,
                    result["pipeline_type"],
                    result["answer"],
                    json.dumps(result["citations"]),
                    json.dumps(result["trace"]),
                    json.dumps(result["metrics"]),
                    json.dumps(result["techniques"]),
                    json.dumps(result["warnings"]),
                    json.dumps(prompt_versions_for_pipeline(result["pipeline_type"])),
                ),
            )
        conn.execute(
            """
            UPDATE runs
            SET traditional_result_id = ?, agentic_result_id = ?, hybrid_result_id = ?
            WHERE id = ?
            """,
            (result_ids.get("traditional"), result_ids.get("agentic"), result_ids.get("hybrid_graph"), run_id),
        )
    for result in results:
        persist_graph_state(run_id, payload.question, payload.selected_sources, result)
    snapshot_run_citations(run_id, results)
    trace_path = write_trace_jsonl(project_id, run_id, results, recommendation)
    final_status = "partial_success" if failed_pipelines else "succeeded"
    record_job_event(job_id, project_id, "persist_outputs", final_status, "Run outputs and trace JSONL persisted.", {"trace_jsonl": str(trace_path)}, run_id)
    run = get_run(run_id)
    write_run_artifacts(run)
    return run


@app.get("/api/projects/{project_id}/runs")
def list_runs(project_id: str) -> list[dict[str, Any]]:
    project_summary(project_id)
    with connect() as conn:
        rows = conn.execute(
            "SELECT * FROM runs WHERE project_id = ? ORDER BY created_at DESC",
            (project_id,),
        ).fetchall()
        runs = []
        for row in rows:
            job = conn.execute(
                "SELECT id, status FROM jobs WHERE run_id = ? ORDER BY created_at DESC LIMIT 1",
                (row["id"],),
            ).fetchone()
            runs.append({
                "id": row["id"],
                "project_id": row["project_id"],
                "knowledge_base_version_id": row["knowledge_base_version_id"],
                "question": row["question"],
                "selected_sources": json.loads(row["selected_sources_json"]),
                "run_mode": row["run_mode"],
                "recommendation": json.loads(row["recommendation_json"]),
                "parent_run_id": row["parent_run_id"],
                "user_feedback": row["user_feedback"],
                "created_at": row["created_at"],
                "job_id": job["id"] if job else None,
                "job_status": job["status"] if job else None,
            })
    return runs


@app.get("/api/runs/{run_id}")
def get_run(run_id: str) -> dict[str, Any]:
    with connect() as conn:
        run = conn.execute("SELECT * FROM runs WHERE id = ?", (run_id,)).fetchone()
        if not run:
            raise HTTPException(status_code=404, detail="Run not found")
        project = conn.execute("SELECT active_version_id FROM projects WHERE id = ?", (run["project_id"],)).fetchone()
        version = conn.execute("SELECT * FROM knowledge_base_versions WHERE id = ?", (run["knowledge_base_version_id"],)).fetchone() if run["knowledge_base_version_id"] else None
        live_citation_chunks = conn.execute(
            """
            SELECT COUNT(*) AS count
            FROM pipeline_results pr, json_each(pr.citations_json) citation
            JOIN chunks c ON c.id = json_extract(citation.value, '$.chunk_id')
            WHERE pr.run_id = ?
            """,
            (run_id,),
        ).fetchone()["count"]
        snapshot_count = conn.execute("SELECT COUNT(*) AS count FROM run_evidence_snapshots WHERE run_id = ?", (run_id,)).fetchone()["count"]
        result_rows = conn.execute(
            "SELECT * FROM pipeline_results WHERE run_id = ?",
            (run_id,),
        ).fetchall()
        job = conn.execute("SELECT id, status FROM jobs WHERE run_id = ? ORDER BY created_at DESC LIMIT 1", (run_id,)).fetchone()
    results = [
        {
            "pipeline_type": row["pipeline_type"],
            "answer": row["answer"],
            "citations": json.loads(row["citations_json"]),
            "trace": json.loads(row["trace_json"]),
            "metrics": json.loads(row["metrics_json"]),
            "techniques": json.loads(row["techniques_json"]),
            "warnings": json.loads(row["warnings_json"]),
            "prompt_versions": json.loads(row["prompt_versions_json"]),
        }
        for row in result_rows
    ]
    run_data = {
        "id": run["id"],
        "project_id": run["project_id"],
        "knowledge_base_version_id": run["knowledge_base_version_id"],
        "question": run["question"],
        "selected_sources": json.loads(run["selected_sources_json"]),
        "run_mode": run["run_mode"],
        "traditional_result_id": run["traditional_result_id"],
        "agentic_result_id": run["agentic_result_id"],
        "hybrid_result_id": run["hybrid_result_id"],
        "recommendation": json.loads(run["recommendation_json"]),
        "comparison": compute_comparison(results),
        "run_config": json.loads(run["run_config_json"]),
        "parent_run_id": run["parent_run_id"],
        "user_feedback": run["user_feedback"],
        "created_at": run["created_at"],
        "results": results,
        "job_id": job["id"] if job else None,
        "job_status": job["status"] if job else None,
        "knowledge_base_version_status": {
            "status": "current" if project and project["active_version_id"] == run["knowledge_base_version_id"] else "archived",
            "version_exists": version is not None,
            "source_evidence_live": live_citation_chunks > 0,
            "snapshot_available": snapshot_count > 0,
        },
    }
    run_root = data_dir() / "projects" / run["project_id"] / "runs" / run["id"]
    if run_root.exists():
        run_data["artifact_paths"] = {
            "traditional": str((run_root / "traditional.json").relative_to(data_dir())),
            "agentic": str((run_root / "agentic.json").relative_to(data_dir())),
            "hybrid_graph": str((run_root / "hybrid_graph.json").relative_to(data_dir())),
            "comparison": str((run_root / "comparison.json").relative_to(data_dir())),
            "trace": str((run_root / "trace.jsonl").relative_to(data_dir())),
        }
    return run_data


@app.post("/api/runs/{run_id}/feedback")
def save_run_feedback(run_id: str, payload: RunFeedbackPatch) -> dict[str, Any]:
    get_run(run_id)
    with connect() as conn:
        conn.execute("UPDATE runs SET user_feedback = ? WHERE id = ?", (payload.user_feedback, run_id))
    return get_run(run_id)


@app.get("/api/runs/{run_id}/citations")
def run_citations(run_id: str) -> list[dict[str, Any]]:
    run = get_run(run_id)
    resolved = []
    for result in run["results"]:
        for citation in result["citations"]:
            item = resolve_citation(chunk_id=citation.get("chunk_id"), label=citation.get("label") or citation.get("citation"), run_id=run_id)
            if item:
                item["pipeline_type"] = result["pipeline_type"]
                resolved.append(item)
    return resolved


def execute_queued_benchmark(job: dict[str, Any]) -> dict[str, Any]:
    payload = job["payload"]
    run_id = payload["run_id"]
    project_id = job["project_id"]
    question = payload["question"]
    selected_sources = payload.get("selected_sources", [])
    if job_cancel_requested(job["id"]):
        return get_run(run_id)
    record_job_event(job["id"], project_id, "create_run", "running", "Benchmark run started by worker.", payload, run_id)
    record_job_event(job["id"], project_id, "start_traditional_pipeline", "running", "Traditional pipeline started.", run_id=run_id)
    traditional = run_pipeline_safely(project_id, question, "traditional", selected_sources)
    if job_cancel_requested(job["id"]):
        record_job_event(job["id"], project_id, "benchmark_cancelled", "cancelled", "Benchmark cancelled after Traditional pipeline checkpoint.", run_id=run_id)
        return get_run(run_id)
    record_job_event(job["id"], project_id, "start_agentic_pipeline", "running", "Agentic pipeline started.", run_id=run_id)
    agentic = run_pipeline_safely(project_id, question, "agentic", selected_sources)
    if job_cancel_requested(job["id"]):
        record_job_event(job["id"], project_id, "benchmark_cancelled", "cancelled", "Benchmark cancelled after Agentic pipeline checkpoint.", run_id=run_id)
        return get_run(run_id)
    record_job_event(job["id"], project_id, "start_hybrid_graph_pipeline", "running", "Hybrid Graph pipeline started.", run_id=run_id)
    hybrid = run_pipeline_safely(project_id, question, "hybrid_graph", selected_sources)
    if job_cancel_requested(job["id"]):
        record_job_event(job["id"], project_id, "benchmark_cancelled", "cancelled", "Benchmark cancelled after Hybrid Graph pipeline checkpoint.", run_id=run_id)
        return get_run(run_id)
    results = [traditional, agentic, hybrid]
    failed_pipelines = [result["pipeline_type"] for result in results if "pipeline_failed" in result.get("warnings", [])]
    record_job_event(job["id"], project_id, "collect_results", "running", "Pipeline results collected.", {"failed_pipelines": failed_pipelines}, run_id)
    recommendation = recommend(results)
    record_job_event(job["id"], project_id, "compute_comparison", "running", "Recommendation and comparison computed.", {"recommended_flow": recommendation["recommended_flow"]}, run_id)
    with connect() as conn:
        result_ids: dict[str, str] = {}
        for result in results:
            result_id = str(uuid.uuid4())
            result_ids[result["pipeline_type"]] = result_id
            conn.execute(
                """
                INSERT INTO pipeline_results
                  (id, run_id, pipeline_type, answer, citations_json, trace_json, metrics_json, techniques_json, warnings_json, prompt_versions_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    result_id,
                    run_id,
                    result["pipeline_type"],
                    result["answer"],
                    json.dumps(result["citations"]),
                    json.dumps(result["trace"]),
                    json.dumps(result["metrics"]),
                    json.dumps(result["techniques"]),
                    json.dumps(result["warnings"]),
                    json.dumps(prompt_versions_for_pipeline(result["pipeline_type"])),
                ),
            )
        conn.execute(
            """
            UPDATE runs
            SET traditional_result_id = ?, agentic_result_id = ?, hybrid_result_id = ?, recommendation_json = ?
            WHERE id = ?
            """,
            (
                result_ids.get("traditional"),
                result_ids.get("agentic"),
                result_ids.get("hybrid_graph"),
                json.dumps(recommendation),
                run_id,
            ),
        )
    for result in results:
        persist_graph_state(run_id, question, selected_sources, result)
    snapshot_run_citations(run_id, results)
    trace_path = write_trace_jsonl(project_id, run_id, results, recommendation)
    final_status = "partial_success" if failed_pipelines else "succeeded"
    record_job_event(job["id"], project_id, "persist_outputs", final_status, "Run outputs and trace JSONL persisted.", {"trace_jsonl": str(trace_path)}, run_id)
    run = get_run(run_id)
    write_run_artifacts(run)
    return run


@app.get("/api/runs/{run_id}/export.json")
def export_json(run_id: str) -> JSONResponse:
    run = get_run(run_id)
    return JSONResponse(
        scrub({
            "schema": "APRAG-Lab.run_export.v1",
            "run": run,
            "privacy": {
                "local_first": True,
                "contains_source_excerpts": True,
                "cloud_provider_used": False,
                "telemetry_enabled": False,
            },
        })
    )


@app.get("/api/runs/{run_id}/events")
def run_events(run_id: str) -> list[dict[str, Any]]:
    get_run(run_id)
    return list_job_events(run_id)


@app.get("/api/runs/{run_id}/events.sse")
def run_events_sse(run_id: str) -> StreamingResponse:
    return StreamingResponse(iter([sse_lines(run_events(run_id))]), media_type="text/event-stream")


@app.websocket("/api/runs/{run_id}/events.ws")
async def run_events_ws(websocket: WebSocket, run_id: str) -> None:
    await websocket.accept()
    try:
        for event in run_events(run_id):
            await websocket.send_json(event)
    except HTTPException:
        await websocket.send_json({"error": "Run not found"})
    await websocket.close()


@app.post("/api/runs/{run_id}/cancel")
def cancel_run(run_id: str) -> dict[str, Any]:
    get_run(run_id)
    try:
        return cancel_run_job(run_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Run job not found") from None


@app.get("/api/runs/{run_id}/graph-states")
def run_graph_states(run_id: str) -> list[dict[str, Any]]:
    get_run(run_id)
    return list_graph_states(run_id)


@app.get("/api/settings/graph-nodes")
def graph_nodes() -> dict[str, Any]:
    return {"nodes": GRAPH_NODES}


@app.get("/api/projects/{project_id}/graph")
def project_graph(project_id: str) -> dict[str, Any]:
    project_summary(project_id)
    with connect() as conn:
        entities = conn.execute("SELECT * FROM entities WHERE project_id = ? ORDER BY source_count DESC, name", (project_id,)).fetchall()
        relationships = conn.execute(
            """
            SELECT r.*, le.name AS from_name, te.name AS to_name, c.citation AS evidence_citation
            FROM relationships r
            JOIN entities le ON le.id = r.from_entity_id
            JOIN entities te ON te.id = r.to_entity_id
            LEFT JOIN chunks c ON c.id = r.evidence_chunk_id
            WHERE r.project_id = ?
            ORDER BY r.confidence DESC
            """,
            (project_id,),
        ).fetchall()
    return {
        "project_id": project_id,
        "nodes": [
            {
                "id": row["id"],
                "name": row["name"],
                "type": row["type"],
                "aliases": json.loads(row["aliases_json"]),
                "source_count": row["source_count"],
                "metadata": json.loads(row["metadata_json"]),
            }
            for row in entities
        ],
        "edges": [
            {
                "id": row["id"],
                "from": row["from_entity_id"],
                "to": row["to_entity_id"],
                "from_name": row["from_name"],
                "to_name": row["to_name"],
                "label": row["label"],
                "relationship_type": row["relationship_type"],
                "confidence": row["confidence"],
                "evidence_chunk_id": row["evidence_chunk_id"],
                "evidence_citation": row["evidence_citation"],
                "metadata": json.loads(row["metadata_json"]),
            }
            for row in relationships
        ],
    }


@app.get("/api/runs/{run_id}/trace.jsonl", response_class=PlainTextResponse)
def run_trace_jsonl(run_id: str) -> str:
    run = get_run(run_id)
    return read_trace_jsonl(run["project_id"], run_id)


@app.get("/api/runs/{run_id}/trace-events")
def run_trace_events(run_id: str) -> dict[str, Any]:
    run = get_run(run_id)
    events = parse_trace_jsonl(read_trace_jsonl(run["project_id"], run_id))
    return {"run_id": run_id, "events": events, "event_count": len(events)}


@app.get("/api/runs/{run_id}/export.md", response_class=PlainTextResponse)
def export_markdown(run_id: str) -> str:
    run = get_run(run_id)
    lines = [f"# APRAG-Lab Run", "", f"Question: {run['question']}", ""]
    lines.append(f"Recommended flow: {run['recommendation']['recommended_flow']}")
    for reason in run["recommendation"]["reasons"]:
        lines.append(f"- {reason}")
    lines.extend(["", "## Comparison", "", f"Best pipeline: {run['comparison']['best_pipeline']}", ""])
    lines.append("Tradeoffs:")
    for note in run["comparison"]["tradeoff_notes"]:
        lines.append(f"- {note}")
    lines.extend(["", "Quality labels:"])
    for name, data in run["comparison"]["pipelines"].items():
        lines.append(f"- {name}: {data['quality_label']} (evidence score: {data['evidence_score']})")
    for result in run["results"]:
        lines.extend(["", f"## {result['pipeline_type']}", "", result["answer"], "", "Metrics:"])
        for key, value in result["metrics"].items():
            lines.append(f"- {key}: {value}")
    return "\n".join(lines)
